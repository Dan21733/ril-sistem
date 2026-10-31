"""Uji: barang titipan (konsinyasi) - pemilik barang, syarat, floor, impor bertujuan, setoran."""
import uuid
from datetime import timedelta
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from pos.models import Item, RefundEvent, RefundEventDetail, TransactionDetail, Transactions

from master.akun.akses import SEMUA_HAK
from master.layanan import excel, inti as layanan, titipan
from master.layanan.inti import Tolak
from master.models import Impor, ItemInfo, Pengaturan, SetoranTitipan, Toko
from master.tests.dasar import Dasar, buat_item


class TitipanTest(Dasar):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.t1 = Toko.objects.get(stok_di_kasir=True)
        cls.brand = Toko.objects.create(kode="DBX", nama="DB X Brand", titip_di=cls.t1, telepon="0812-1111")
        buat_item("Kloset Brand", "CW1", floor=0, jual=67500, modal=None, kategori="Sanitary", brand="X")
        buat_item("Kloset Brand", "CW2", floor=0, jual=100000, modal=None, kategori="Sanitary", brand="X")
        layanan.sinkron_info()
        ItemInfo.objects.filter(item_name="Kloset Brand").update(toko=cls.brand)

    def setUp(self):
        self.client.force_login(self.admin)

    def _jual(self, nama, varian, qty, harga, waktu=None):
        sid = uuid.uuid4()
        Transactions.objects.create(sync_id=sid, transaction_id=f"T01-{uuid.uuid4().hex[:6]}", worker_id="ID001",
                                    transaction_time=waktu or timezone.now(), item_amount=qty * harga, shipping_amount=0,
                                    total_amount=qty * harga, payment_method="cash", payment_amount=qty * harga, change_amount=0)
        TransactionDetail.objects.create(sync_id=sid, item_name=nama, variant=varian, quantity=qty, unit_price=harga,
                                         discount=0, subtotal=qty * harga)
        return sid

    # ---- syarat & floor ----

    def test_floor_dari_diskon_maks_dibulatkan_ke_atas(self):
        info = self.info("Kloset Brand", "CW1")
        titipan.simpan_syarat(self.admin, info, potongan="15", diskon_maks="10")
        self.assertEqual(Item.objects.get(item_name="Kloset Brand", variant="CW1").min_price, 60800)  # 60.750 -> 60.800
        Pengaturan.objects.filter(pk=1).update(pembulatan_diskon=1000)
        titipan.segarkan_floor(self.admin, info)
        self.assertEqual(Item.objects.get(item_name="Kloset Brand", variant="CW1").min_price, 61000)

    def test_floor_titipan_tidak_bisa_diketik_dan_ikut_harga_jual(self):
        info = self.info("Kloset Brand", "CW1")
        titipan.simpan_syarat(self.admin, info, potongan="15", diskon_maks="10")
        with self.assertRaises(Tolak):
            layanan.ubah_varian(self.admin, info, "min_price", "50000", hak=SEMUA_HAK)
        layanan.ubah_varian(self.admin, info, "max_price", "80000", hak=SEMUA_HAK)
        self.assertEqual(Item.objects.get(item_name="Kloset Brand", variant="CW1").min_price, 72000)

    def test_syarat_hanya_untuk_barang_titipan_dan_tanggal_tidak_di_masa_depan(self):
        with self.assertRaises(Tolak):
            titipan.simpan_syarat(self.admin, self.info("Palu", "Kecil"), potongan="10", diskon_maks="5")
        with self.assertRaises(Tolak):
            titipan.simpan_syarat(self.admin, self.info("Kloset Brand", "CW1"), potongan="10", diskon_maks="5",
                                  berlaku_mulai=timezone.localdate() + timedelta(days=1))

    def test_panel_menyimpan_syarat_semua_varian(self):
        item = ItemInfo.objects.get(item_name="Kloset Brand")
        cw1, cw2 = self.info("Kloset Brand", "CW1"), self.info("Kloset Brand", "CW2")
        panel = self.client.get(reverse("master:item_panel", args=[item.id]))
        self.assertContains(panel, "Syarat titipan")
        self.client.post(reverse("master:item", args=[item.id]), {
            "aksi": "syarat_titipan", f"pot_{cw1.id}": "12", f"maks_{cw1.id}": "5",
            f"pot_{cw2.id}": "12,5", f"maks_{cw2.id}": "", "berlaku_mulai": timezone.localdate().isoformat()})
        kini = titipan.syarat_kini([cw1.id, cw2.id])
        self.assertEqual((str(kini[cw1.id].potongan_toko), str(kini[cw1.id].diskon_maks)), ("12.00", "5.00"))
        self.assertEqual((str(kini[cw2.id].potongan_toko), kini[cw2.id].diskon_maks), ("12.50", None))
        # Barang toko sendiri tidak punya bagian ini.
        self.assertNotContains(self.client.get(reverse("master:item_panel", args=[ItemInfo.objects.get(item_name="Palu").id])),
                               "Syarat titipan")

    # ---- daftar barang & barang baru per toko ----

    def test_daftar_barang_mengikuti_toko_aktif(self):
        isi = self.client.get(reverse("master:barang")).content.decode()
        self.assertIn("Palu", isi)
        self.assertNotIn("Kloset Brand", isi)
        self.client.post(reverse("master:pilih_toko"), {"toko": self.brand.pk, "next": reverse("master:barang")})
        isi = self.client.get(reverse("master:barang")).content.decode()
        self.assertIn("Kloset Brand", isi)
        self.assertNotIn("Palu", isi)
        self.assertIn("DB X Brand", isi)

    def test_barang_baru_saat_toko_titipan_aktif_menjadi_milik_titipan(self):
        import json
        self.client.post(reverse("master:pilih_toko"), {"toko": self.brand.pk, "next": "/"})
        r = self.client.post(reverse("master:barang_baru"), json.dumps({
            "item_name": "Wastafel Brand", "category": "Sanitary",
            "varian": [{"variant": "", "min_price": "", "max_price": "250000", "in_stock": "0"}]}),
            content_type="application/json").json()
        self.assertTrue(r["ok"])
        self.assertEqual(ItemInfo.objects.get(item_name="Wastafel Brand").toko, self.brand)

    # ---- impor ----

    def _berkas(self, baris, kepala=("item_name", "variant", "category", "min_price", "max_price",
                                     "potongan_toko_persen", "diskon_maks_persen")):
        wb = Workbook()
        ws = wb.active
        ws.title = "item"
        ws.append(list(kepala))
        for b in baris:
            ws.append(list(b))
        isi = BytesIO()
        wb.save(isi)
        isi.seek(0)
        return isi

    def test_impor_ke_toko_titipan(self):
        berkas = self._berkas([("Shower Brand", "S1", "Sanitary", None, 200000, 15, 10),
                               ("Kloset Brand", "CW1", "Sanitary", None, 67500, 12, 5),
                               ("Palu", "Kecil", "Paku", 1000, 1500, 10, 5)])
        isi = excel.baca_berkas(berkas, SEMUA_HAK, toko=self.brand)
        self.assertEqual(isi["toko_nama"], "DB X Brand")
        self.assertEqual([b["item_name"] for b in isi["baru"]], ["Shower Brand"])
        self.assertTrue(any("Palu" in g["pesan"] and "milik" in g["pesan"] for g in isi["galat"]))
        self.assertEqual({x["item_name"] for x in isi["syarat"]}, {"Shower Brand", "Kloset Brand"})
        impor = Impor.objects.create(pengguna=self.admin, nama_berkas="brand.xlsx", isi=isi)
        hasil = excel.terapkan(impor, self.admin, SEMUA_HAK, izinkan_baru=True)
        self.assertEqual((hasil["baru"], hasil["syarat"]), (1, 2))
        self.assertEqual(ItemInfo.objects.get(item_name="Shower Brand").toko, self.brand)
        self.assertEqual(Item.objects.get(item_name="Shower Brand").min_price, 180000)
        self.assertEqual(Item.objects.get(item_name="Kloset Brand", variant="CW1").min_price, 64200)  # 64.125 -> 64.200

    def test_impor_ke_toko_sendiri_menolak_barang_titipan(self):
        isi = excel.baca_berkas(self._berkas([("Kloset Brand", "CW1", "Sanitary", 1, 67500, None, None)]),
                                SEMUA_HAK, toko=self.t1)
        self.assertTrue(isi["galat"])
        self.assertFalse(isi["perubahan"] or isi["syarat"])
        self.assertIn("diskon_maks_persen", isi["kolom_diabaikan"])

    def test_export_toko_titipan_membawa_kolom_syarat(self):
        from openpyxl import load_workbook
        titipan.simpan_syarat(self.admin, self.info("Kloset Brand", "CW1"), potongan="15", diskon_maks="10")
        wb = load_workbook(BytesIO(excel.buat_workbook(titipan.saring_barang(Item.objects.all(), self.brand),
                                                       SEMUA_HAK, toko=self.brand)))
        baris = list(wb["item"].iter_rows(values_only=True))
        self.assertIn("potongan_toko_persen", baris[0])
        self.assertEqual({r[1] for r in baris[1:]}, {"Kloset Brand"})

    def test_layar_impor_wajib_tujuan(self):
        berkas = self._berkas([("Shower Brand", "S1", "Sanitary", None, 200000, 15, 10)])
        self.client.post(reverse("master:impor"), {"berkas": SimpleUploadedFile("b.xlsx", berkas.getvalue())})
        self.assertFalse(Impor.objects.exists())
        self.client.post(reverse("master:impor"), {"berkas": SimpleUploadedFile("b.xlsx", berkas.getvalue()),
                                                  "toko": self.brand.pk})
        impor = Impor.objects.get()
        self.assertContains(self.client.get(reverse("master:impor_pratinjau", args=[impor.id])), "Untuk toko: <b>DB X Brand</b>")

    # ---- setoran & laporan ----

    def test_setoran_dikurangi_retur_dan_potongan_per_tanggal(self):
        cw1 = self.info("Kloset Brand", "CW1")
        kemarin = timezone.localdate() - timedelta(days=1)
        titipan.simpan_syarat(self.admin, cw1, potongan="10", diskon_maks="5", berlaku_mulai=kemarin)
        titipan.simpan_syarat(self.admin, cw1, potongan="20", diskon_maks="5")
        self._jual("Kloset Brand", "CW1", 1, 60000, waktu=timezone.now() - timedelta(days=1))   # potongan 10%
        sid = self._jual("Kloset Brand", "CW1", 2, 65000)                                       # potongan 20%
        self._jual("Palu", "Kecil", 3, 1500)                                                    # bukan titipan
        e = RefundEvent.objects.create(refund_id=uuid.uuid4(), sync_id=sid, transaction_id="x", event_type="return",
                                       reason="rusak", refund_method="cash", refund_amount=65000,
                                       approver_worker_id="ID001", device_code="T01", created_at=timezone.now())
        RefundEventDetail.objects.create(refund_id=e.refund_id, item_name="Kloset Brand", variant="CW1", quantity=1,
                                         unit_price=65000, subtotal=65000)
        h = titipan.hitung(self.brand, kemarin, timezone.localdate())
        t = h["total"]
        self.assertEqual((t["jual"], t["retur"]), (190000, 65000))
        self.assertEqual(t["potongan"], 6000 + 26000 - 13000)
        self.assertEqual(t["disetor"], 190000 - 65000 - 19000)

        # Laporan penjualan: laba barang titipan = potongan toko.
        r = self.client.get(reverse("master:laporan") + "?periode=hari")
        self.assertEqual(r.status_code, 200)

        s = titipan.tandai_disetor(self.admin, self.brand, kemarin, timezone.localdate(), "transfer")
        self.assertEqual(s.disetor, 106000)
        self.assertEqual(titipan.awal_periode(self.brand), timezone.localdate() + timedelta(days=1))
        with self.assertRaises(Tolak):
            titipan.tandai_disetor(self.admin, self.brand, kemarin, timezone.localdate())

    def test_setoran_ditolak_kalau_potongan_belum_diisi(self):
        self._jual("Kloset Brand", "CW2", 1, 100000)
        h = titipan.hitung(self.brand, timezone.localdate(), timezone.localdate())
        self.assertEqual(h["tanpa_syarat"], 1)
        with self.assertRaises(Tolak):
            titipan.tandai_disetor(self.admin, self.brand, timezone.localdate(), timezone.localdate())

    def test_laba_laporan_barang_titipan_adalah_potongan(self):
        from datetime import datetime, time
        from master.layanan import laporan
        titipan.simpan_syarat(self.admin, self.info("Kloset Brand", "CW1"), potongan="15", diskon_maks="5")
        self._jual("Kloset Brand", "CW1", 2, 65000)
        zona = timezone.get_current_timezone()
        awal = timezone.make_aware(datetime.combine(timezone.localdate(), time.min), zona)
        _, baris = laporan.ambil(awal, awal + timedelta(days=1))
        b = next(x for x in baris if x["item_name"] == "Kloset Brand")
        self.assertEqual((b["laba"], b["modal"], b["titipan"]), (19500, 110500, "DB X Brand"))

    def test_halaman_setoran_dan_menu(self):
        titipan.simpan_syarat(self.admin, self.info("Kloset Brand", "CW1"), potongan="15", diskon_maks="5")
        self._jual("Kloset Brand", "CW1", 1, 65000)
        r = self.client.get(reverse("master:titipan"), follow=True)
        self.assertContains(r, "Setoran DB X Brand")
        self.assertContains(r, "wa.me/628121111")
        self.assertContains(r, "55.250")
        self.assertContains(self.client.get(reverse("master:barang")), "Setoran titipan")
        x = self.client.get(reverse("master:titipan_toko", args=[self.brand.pk]) + "?excel=1")
        self.assertEqual(x["Content-Type"].split(";")[0], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        hari = timezone.localdate().isoformat()
        self.client.post(reverse("master:titipan_toko", args=[self.brand.pk]), {"dari": hari, "sampai": hari, "catatan": "tf"})
        s = SetoranTitipan.objects.get()
        self.assertContains(self.client.get(reverse("master:titipan_setoran", args=[s.id])), "55.250")
        # Akun tanpa hak lihat modal tidak bisa membuka setoran (potongan toko rahasia).
        self.client.force_login(self.gudang)
        self.assertNotEqual(self.client.get(reverse("master:titipan_toko", args=[self.brand.pk])).status_code, 200)

    def test_form_toko_titipan(self):
        from master.layanan import toko as svc_toko
        t = svc_toko.simpan(self.admin, {"kode": "BR2", "nama": "Brand Dua", "titip_di": str(self.t1.pk)})
        self.assertEqual(t.titip_di, self.t1)
        with self.assertRaises(Tolak):
            svc_toko.simpan(self.admin, {"kode": "T1", "nama": self.t1.nama, "titip_di": str(self.brand.pk)}, self.t1)
        with self.assertRaises(Tolak):   # masih punya barang titipan
            svc_toko.simpan(self.admin, {"kode": "DBX", "nama": "DB X Brand", "titip_di": ""}, self.brand)
        svc_toko.simpan_pembulatan(self.admin, "500")
        self.assertEqual(Pengaturan.ambil().pembulatan_diskon, 500)
