"""Uji: Barang, harga, modal, import/export Excel, halaman umum, sampah."""
import json
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from pos.models import Item, TransactionDetail

from master.layanan import excel
from master.layanan import inti as layanan
from master.layanan import masalah
from master.akun.akses import PERAN, SEMUA_HAK
from master.layanan.inti import Tolak
from master.models import Impor, ItemInfo, Log, VarianInfo
from master.tests.dasar import Dasar, buat_item, hasil_batch


class AngkaTest(TestCase):
    def test_bentuk_angka(self):
        self.assertEqual(layanan.baca_angka("1.450.000", False), 1450000)
        self.assertEqual(layanan.baca_angka("Rp 17.000", False), 17000)
        self.assertEqual(layanan.baca_angka("17,000", False), 17000)
        self.assertEqual(layanan.baca_angka(17000.0, False), 17000)
        self.assertEqual(layanan.baca_angka("17000.00", False), 17000)
        self.assertIsNone(layanan.baca_angka("", True))

    def test_angka_ditolak(self):
        for buruk in ["abc", "-5", "1.5", 12.5, "1.45.000"]:
            with self.assertRaises(Tolak, msg=buruk):
                layanan.baca_angka(buruk, False)
        with self.assertRaises(Tolak):
            layanan.baca_angka("", False)

    def test_alias_dirapikan(self):
        self.assertEqual(layanan.rapikan_alias(" Martil,  hammer ; MARTIL,,penokok "), "martil, hammer, penokok")


class UbahTest(Dasar):
    def test_ubah_harga_tercatat(self):
        info = self.info("Palu", "Kecil")
        nilai, berubah = layanan.ubah_varian(self.admin, info, "max_price", "2.000", hak=SEMUA_HAK)
        self.assertEqual((nilai, berubah), (2000, True))
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").max_price, 2000)
        log = Log.objects.get(kolom="max_price")
        self.assertEqual((log.lama, log.baru, log.nama_pengguna, log.rahasia), ("1500", "2000", "admin", False))

    def test_nilai_sama_tidak_dicatat(self):
        _, berubah = layanan.ubah_varian(self.admin, self.info("Palu", "Kecil"), "max_price", "1500", hak=SEMUA_HAK)
        self.assertFalse(berubah)
        self.assertEqual(Log.objects.count(), 0)

    def test_finance_tidak_boleh_ubah_modal(self):
        with self.assertRaises(Tolak):
            layanan.ubah_varian(self.finance, self.info("Palu", "Kecil"), "cost_amount", "900", hak=PERAN["Finance"])
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").cost_amount, 800)

    def test_gudang_tidak_boleh_ubah_harga(self):
        with self.assertRaises(Tolak):
            layanan.ubah_varian(self.gudang, self.info("Palu", "Kecil"), "min_price", "1", hak=PERAN["Gudang"])

    def test_log_modal_ditandai_rahasia(self):
        layanan.ubah_varian(self.admin, self.info("Palu", "Kecil"), "cost_amount", "900", hak=SEMUA_HAK)
        self.assertTrue(Log.objects.get(kolom="cost_amount").rahasia)

    def test_rak_harus_terdaftar(self):
        info = self.info("Palu", "Kecil")
        with self.assertRaises(Tolak):
            layanan.ubah_varian(self.admin, info, "kode_rak", "ZZ-99", hak=SEMUA_HAK)
        layanan.ubah_varian(self.admin, info, "kode_rak", "TK-01", hak=SEMUA_HAK)
        self.assertEqual(self.info("Palu", "Kecil").rak_id, "TK-01")

    def test_alias_item_berlaku_semua_varian(self):
        layanan.ubah_alias_item(self.admin, ItemInfo.objects.get(item_name="Palu"), "Martil, penokok", hak=SEMUA_HAK)
        self.assertEqual(ItemInfo.objects.get(item_name="Palu").alias, "martil, penokok")
        self.assertEqual(Log.objects.get(kolom="alias").variant, "*")


class GantiNamaTest(Dasar):
    def test_ganti_nama_item_membawa_nota(self):
        item_info = ItemInfo.objects.get(item_name="Palu")
        id_lama = self.info("Palu", "Kecil").id
        layanan.ganti_nama_item(self.admin, item_info, "Palu Kambing", hak=SEMUA_HAK)
        self.assertFalse(Item.objects.filter(item_name="Palu").exists())
        self.assertEqual(Item.objects.filter(item_name="Palu Kambing").count(), 2)
        self.assertEqual(TransactionDetail.objects.get(sync_id=self.nota.sync_id).item_name, "Palu Kambing")
        # ID tetap tidak berubah walau nama diganti.
        self.assertEqual(self.info("Palu Kambing", "Kecil").id, id_lama)
        self.assertEqual(Item.objects.get(item_name="Palu Kambing", variant="Kecil").cost_amount, 800)

    def test_ganti_nama_ke_nama_terpakai_ditolak(self):
        with self.assertRaises(Tolak):
            layanan.ganti_nama_item(self.admin, ItemInfo.objects.get(item_name="Palu"), "Kikir", hak=SEMUA_HAK)
        self.assertEqual(Item.objects.filter(item_name="Palu").count(), 2)

    def test_ganti_nama_varian(self):
        info = self.info("Palu", "Kecil")
        layanan.ganti_nama_varian(self.admin, info, "8 Oz", hak=SEMUA_HAK)
        self.assertTrue(Item.objects.filter(item_name="Palu", variant="8 Oz").exists())
        self.assertEqual(TransactionDetail.objects.get(sync_id=self.nota.sync_id).variant, "8 Oz")
        with self.assertRaises(Tolak):
            layanan.ganti_nama_varian(self.admin, info, "Besar", hak=SEMUA_HAK)

    def test_finance_tidak_boleh_ganti_nama(self):
        with self.assertRaises(Tolak):
            layanan.ganti_nama_item(self.finance, ItemInfo.objects.get(item_name="Palu"), "X", hak=PERAN["Finance"])


class MasalahTest(Dasar):
    def test_saring_dan_periksa_sejalan(self):
        hak = SEMUA_HAK
        semua = list(Item.objects.values("item_name", "variant", "min_price", "max_price", "cost_amount"))
        alias = dict(ItemInfo.objects.values_list("item_name", "alias"))
        rak = {(v.item_name, v.variant): v.rak_id for v in VarianInfo.objects.all()}
        for b in semua:
            b["alias"] = alias[b["item_name"]]
            b["kode_rak"] = rak[(b["item_name"], b["variant"])]
        for aturan in masalah.aturan_untuk(hak):
            dari_db = set(aturan["saring"](Item.objects.all()).values_list("item_name", "variant"))
            dari_py = {(b["item_name"], b["variant"]) for b in semua if aturan["periksa"](b)}
            self.assertEqual(dari_db, dari_py, aturan["kunci"])

    def test_jumlah_yang_diharapkan(self):
        hitung = {a["kunci"]: a["saring"](Item.objects.all()).count() for a in masalah.ATURAN}
        self.assertEqual(hitung["floor_terbalik"], 2)       # Palu Besar, Oracle
        self.assertEqual(hitung["harga_nol"], 1)            # Kikir
        self.assertEqual(hitung["jual_dibawah_modal"], 1)   # Oracle
        self.assertEqual(hitung["modal_kosong"], 1)         # Kikir
        self.assertEqual(hitung["alias_kosong"], 4)

    def test_aturan_modal_disembunyikan_dari_finance(self):
        kunci = {a["kunci"] for a in masalah.aturan_untuk(PERAN["Finance"])}
        self.assertNotIn("jual_dibawah_modal", kunci)
        self.assertNotIn("modal_kosong", kunci)


class ExcelTest(Dasar):
    def _berkas(self, wb):
        isi = BytesIO()
        wb.save(isi)
        isi.seek(0)
        return isi

    def test_export_tanpa_modal_untuk_finance(self):
        wb = load_workbook(BytesIO(excel.buat_workbook(Item.objects.all(), PERAN["Finance"])))
        kepala = [c.value for c in wb["item"][1]]
        self.assertNotIn("cost_amount", kepala)
        wb = load_workbook(BytesIO(excel.buat_workbook(Item.objects.all(), SEMUA_HAK)))
        self.assertIn("cost_amount", [c.value for c in wb["item"][1]])

    def test_export_lalu_impor_tanpa_ubah_tidak_ada_perubahan(self):
        isi = excel.baca_berkas(BytesIO(excel.buat_workbook(Item.objects.all(), SEMUA_HAK)), SEMUA_HAK)
        self.assertEqual((isi["perubahan"], isi["alias"], isi["baru"], isi["galat"]), ([], [], [], []))

    def test_impor_ubah_harga_alias_dan_rak(self):
        wb = load_workbook(BytesIO(excel.buat_workbook(Item.objects.all(), SEMUA_HAK)))
        ws = wb["item"]
        kepala = [c.value for c in ws[1]]
        for baris in ws.iter_rows(min_row=2):
            sel = dict(zip(kepala, baris))
            if (sel["item_name"].value, sel["variant"].value) == ("Palu", "Besar"):
                sel["min_price"].value = 2000
                sel["kode_rak"].value = "GB-01"
            if sel["item_name"].value == "Palu":
                sel["alias"].value = "martil, penokok"
        wb["rak"].append(["GB-01", "Rak Lembaran", "Gudang B", ""])
        isi = excel.baca_berkas(self._berkas(wb), SEMUA_HAK)
        self.assertEqual(len(isi["perubahan"]), 2)
        self.assertEqual(len(isi["alias"]), 1)
        self.assertEqual(isi["rak"][0]["jenis"], "baru")
        self.assertEqual(isi["galat"], [])
        # Belum ada yang berubah sebelum diterapkan.
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Besar").min_price, 3000)

        impor = Impor.objects.create(pengguna=self.admin, nama_berkas="uji.xlsx", isi=isi)
        hasil = excel.terapkan(impor, self.admin, SEMUA_HAK)
        self.assertEqual((hasil["diubah"], hasil["alias"], hasil["rak"]), (2, 1, 1))
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Besar").min_price, 2000)
        self.assertEqual(self.info("Palu", "Besar").rak_id, "GB-01")
        self.assertEqual(ItemInfo.objects.get(item_name="Palu").alias, "martil, penokok")
        self.assertEqual(Log.objects.filter(sumber="impor", batch=hasil_batch(impor)).count(), 4)

    def test_impor_angka_buruk_dan_rak_asing_dilewati(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "item"
        ws.append(["item_name", "variant", "min_price", "max_price", "kode_rak"])
        ws.append(["Palu", "Kecil", "seribu", 1600, "ZZ-99"])
        isi = excel.baca_berkas(self._berkas(wb), SEMUA_HAK)
        self.assertEqual([p["kolom"] for p in isi["perubahan"]], ["max_price"])
        self.assertEqual(len(isi["galat"]), 2)

    def test_impor_barang_baru_perlu_izin(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "item"
        ws.append(["item_name", "variant", "category", "min_price", "max_price"])
        ws.append(["Obeng", None, "Alat", 5000, 7000])
        isi = excel.baca_berkas(self._berkas(wb), SEMUA_HAK)
        self.assertEqual(len(isi["baru"]), 1)
        impor = Impor.objects.create(pengguna=self.admin, nama_berkas="uji.xlsx", isi=isi)
        excel.terapkan(impor, self.admin, SEMUA_HAK)
        self.assertFalse(Item.objects.filter(item_name="Obeng").exists())
        impor2 = Impor.objects.create(pengguna=self.admin, nama_berkas="uji.xlsx", isi=isi)
        excel.terapkan(impor2, self.admin, SEMUA_HAK, izinkan_baru=True)
        self.assertEqual(Item.objects.get(item_name="Obeng").variant, "N/A")
        self.assertTrue(VarianInfo.objects.filter(item_name="Obeng", variant="N/A").exists())

    def test_finance_impor_tidak_menyentuh_modal(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "item"
        ws.append(["item_name", "variant", "cost_amount", "max_price"])
        ws.append(["Palu", "Kecil", 1, 1700])
        isi = excel.baca_berkas(self._berkas(wb), PERAN["Finance"])
        self.assertEqual([p["kolom"] for p in isi["perubahan"]], ["max_price"])
        self.assertEqual(isi["kolom_tanpa_hak"], ["cost_amount"])


class HalamanTest(Dasar):
    def test_wajib_login(self):
        for nama in ["dashboard", "barang", "kategori", "brand", "rak", "impor", "log", "pengguna"]:
            respons = self.client.get(reverse(f"master:{nama}"))
            self.assertEqual(respons.status_code, 302, nama)
            self.assertIn("/master/login/", respons["Location"])

    def test_semua_halaman_terbuka_untuk_admin(self):
        self.client.force_login(self.admin)
        for nama in ["dashboard", "barang", "kategori", "brand", "rak", "impor", "log", "pengguna"]:
            self.assertEqual(self.client.get(reverse(f"master:{nama}")).status_code, 200, nama)
        self.assertEqual(self.client.get(reverse("master:barang_baru") + "?panel=1").status_code, 200)
        item_id = ItemInfo.objects.get(item_name="Palu").id
        self.assertEqual(self.client.get(reverse("master:item", args=[item_id])).status_code, 200)

    def test_finance_tidak_melihat_modal_di_mana_pun(self):
        Item.objects.filter(item_name="Palu", variant="Kecil").update(cost_amount=987654)
        self.client.force_login(self.finance)
        item_id = ItemInfo.objects.get(item_name="Palu").id
        for alamat in [reverse("master:barang"), reverse("master:item", args=[item_id]), reverse("master:dashboard")]:
            isi = self.client.get(alamat).content.decode()
            self.assertNotIn("987654", isi)
            self.assertNotIn("987.654", isi)
            self.assertNotIn("Modal", isi.replace("Harga &amp; Modal", ""))
        respons = self.client.get(reverse("master:barang") + "?export=1")
        wb = load_workbook(BytesIO(respons.content))
        self.assertNotIn("cost_amount", [c.value for c in wb["item"][1]])
        self.assertEqual(self.client.get(reverse("master:pengguna")).status_code, 403)
        self.assertEqual(self.client.get(reverse("master:impor")).status_code, 403)

    def test_finance_tidak_melihat_modal_di_log(self):
        layanan.ubah_varian(self.admin, self.info("Palu", "Kecil"), "cost_amount", "987654", hak=SEMUA_HAK)
        self.client.force_login(self.finance)
        isi = self.client.get(reverse("master:log")).content.decode()
        self.assertNotIn("987654", isi)
        self.assertIn("disembunyikan", isi)
        # Mencari nilai modal tidak boleh membocorkan baris yang cocok.
        self.assertIn("0 catatan", self.client.get(reverse("master:log") + "?q=987654").content.decode())

    def test_edit_sel(self):
        self.client.force_login(self.admin)
        info = self.info("Palu", "Besar")
        respons = self.client.post(reverse("master:barang_ubah"), json.dumps(
            {"id": info.id, "kolom": "min_price", "nilai": "2.000"}), content_type="application/json")
        data = respons.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["nilai"], 2000)
        self.assertNotIn("floor_terbalik", [m["kunci"] for m in data["masalah"]])

    def test_edit_sel_ditolak_untuk_gudang(self):
        self.client.force_login(self.gudang)
        respons = self.client.post(reverse("master:barang_ubah"), json.dumps(
            {"id": self.info("Palu", "Kecil").id, "kolom": "max_price", "nilai": "1"}), content_type="application/json")
        self.assertEqual(respons.status_code, 403)
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").max_price, 1500)

    def test_filter_dan_cari_alias(self):
        layanan.ubah_alias_item(self.admin, ItemInfo.objects.get(item_name="Palu"), "martil", hak=SEMUA_HAK)
        self.client.force_login(self.admin)
        isi = self.client.get(reverse("master:barang") + "?q=martil").content.decode()
        self.assertIn("Palu", isi)
        self.assertNotIn("Kikir", isi)
        isi = self.client.get(reverse("master:barang") + "?masalah=harga_nol").content.decode()
        self.assertIn("Kikir", isi)
        self.assertNotIn("Oracle", isi)

    def test_unggah_impor_lewat_layar(self):
        self.client.force_login(self.admin)
        isi = excel.buat_workbook(Item.objects.all(), SEMUA_HAK)
        berkas = SimpleUploadedFile("barang.xlsx", isi)
        # Tujuan toko wajib dipilih: tanpa itu berkas tidak dibaca.
        self.client.post(reverse("master:impor"), {"berkas": SimpleUploadedFile("barang.xlsx", isi)})
        self.assertFalse(Impor.objects.exists())
        from master.models import Toko
        respons = self.client.post(reverse("master:impor"), {"berkas": berkas, "toko": Toko.objects.get(stok_di_kasir=True).pk})
        self.assertEqual(respons.status_code, 302)
        impor = Impor.objects.get()
        self.assertEqual(self.client.get(reverse("master:impor_pratinjau", args=[impor.id])).status_code, 200)

    def test_tambah_pengguna(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("master:pengguna"), {"aksi": "tambah", "username": "Ibu", "password": "kata-sandi-1", "templat": "pemilik"})
        akun = get_user_model().objects.get(username="ibu")
        self.assertTrue(akun.check_password("kata-sandi-1"))
        from master.akun.akses import hak_pengguna
        self.assertEqual(hak_pengguna(akun), set(SEMUA_HAK))
        self.assertEqual(self.client.get(reverse("master:pengguna")).status_code, 200)


class DaftarPerBarangTest(Dasar):
    def test_urutan_alami(self):
        from master.views.utama import kunci_alami
        acak = ["Cw633 + 505", "CW 53 + 01R", "CW 415 + 505", "Cw420 + 505", "CW 427 + 505", "cw 9"]
        self.assertEqual(
            sorted(acak, key=kunci_alami),
            ["cw 9", "CW 53 + 01R", "CW 415 + 505", "Cw420 + 505", "CW 427 + 505", "Cw633 + 505"],
        )

    def test_satu_baris_per_barang(self):
        self.client.force_login(self.admin)
        respons = self.client.get(reverse("master:barang"))
        daftar = respons.context["daftar"]
        self.assertEqual([d["item_name"] for d in daftar], ["Kikir", "Oracle", "Palu"])
        palu = daftar[2]
        self.assertEqual((palu["total"], palu["harga_min"], palu["harga_max"]), (2, 1500, 2500))
        self.assertEqual(respons.context["jumlah_varian"], 4)

    def test_urut_nama_tidak_peduli_huruf_besar(self):
        buat_item("amplas")
        buat_item("Zinc")
        layanan.sinkron_info()
        self.client.force_login(self.admin)
        nama = [d["item_name"] for d in self.client.get(reverse("master:barang")).context["daftar"]]
        self.assertEqual(nama, ["amplas", "Kikir", "Oracle", "Palu", "Zinc"])

    def test_panel_dan_kembali_ke_daftar(self):
        self.client.force_login(self.admin)
        item_id = ItemInfo.objects.get(item_name="Palu").id
        isi = self.client.get(reverse("master:item_panel", args=[item_id]) + "?kembali=/master/barang/%3Fbuka%3D1").content.decode()
        self.assertNotIn("<html", isi)
        self.assertIn("data-tutup-panel", isi)
        self.assertLess(isi.index(">Besar<"), isi.index(">Kecil<"))
        respons = self.client.post(reverse("master:item", args=[item_id]), {
            "aksi": "simpan_item", "alias": "martil", "category": "Paku", "brand": "ABC",
            "kembali": "/master/barang/?q=palu&buka=%d" % item_id,
        })
        self.assertEqual(respons["Location"], "/master/barang/?q=palu&buka=%d" % item_id)
        # Alamat kembali ke situs lain tidak diikuti.
        respons = self.client.post(reverse("master:item", args=[item_id]), {
            "aksi": "simpan_item", "alias": "martil", "category": "Paku", "brand": "ABC",
            "kembali": "https://contoh.jahat/",
        })
        self.assertEqual(respons["Location"], reverse("master:item", args=[item_id]))


class HargaTest(Dasar):
    def test_hitung_massal(self):
        self.assertEqual(layanan.hitung_massal(1000, "persen", 5, 0), 1050)
        self.assertEqual(layanan.hitung_massal(1040, "persen", 5, 500), 1500)    # 1092 -> ke atas
        self.assertEqual(layanan.hitung_massal(1000, "rupiah", -300, 0), 700)
        self.assertEqual(layanan.hitung_massal(200, "rupiah", -300, 0), 0)       # tidak pernah negatif
        self.assertEqual(layanan.hitung_massal(None, "persen", 5, 0), None)      # kosong tidak dihitung
        self.assertEqual(layanan.hitung_massal(None, "tetapkan", 900, 0), 900)

    def test_parameter_ditolak(self):
        for data in [
            {"kolom": "in_stock", "cara": "persen", "angka": "5"},
            {"kolom": "max_price", "cara": "persen", "angka": "abc"},
            {"kolom": "max_price", "cara": "persen", "angka": "0"},
            {"kolom": "max_price", "cara": "persen", "angka": "900"},
            {"kolom": "max_price", "cara": "tetapkan", "angka": "-5"},
            {"kolom": "max_price", "cara": "persen", "angka": "5", "bulat": "7"},
        ]:
            with self.assertRaises(Tolak, msg=data):
                layanan.baca_parameter_massal(data, SEMUA_HAK)
        with self.assertRaises(Tolak):
            layanan.baca_parameter_massal({"kolom": "cost_amount", "cara": "persen", "angka": "5"}, PERAN["Finance"])

    def test_pratinjau_tidak_mengubah_lalu_terapkan_mengubah(self):
        self.client.force_login(self.admin)
        alamat = reverse("master:harga") + "?q=palu"
        data = {"kolom": "max_price", "cara": "persen", "angka": "10", "bulat": "500", "catatan": "supplier naik"}
        respons = self.client.post(alamat, {**data, "aksi": "pratinjau"})
        self.assertEqual(respons.context["pratinjau"]["jumlah"], 2)
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").max_price, 1500)
        self.client.post(alamat, {**data, "aksi": "terapkan"})
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").max_price, 2000)   # 1650 -> 2000
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Besar").max_price, 3000)   # 2750 -> 3000
        self.assertEqual(Item.objects.get(item_name="Kikir").max_price, 0)                      # di luar filter
        log = Log.objects.filter(sumber="massal")
        self.assertEqual(log.count(), 2)
        self.assertEqual({l.catatan for l in log}, {"supplier naik"})
        self.assertEqual(len({l.batch for l in log}), 1)

    def test_pratinjau_memperingatkan_floor_terbalik(self):
        self.client.force_login(self.admin)
        respons = self.client.post(reverse("master:harga") + "?q=palu+kecil", {
            "kolom": "min_price", "cara": "tetapkan", "angka": "9000", "aksi": "pratinjau"})
        self.assertEqual(respons.context["pratinjau"]["terbalik"], 1)

    def test_gudang_tidak_bisa_membuka(self):
        self.client.force_login(self.gudang)
        self.assertRedirects(self.client.get(reverse("master:harga")), reverse("gudang:beranda"))

    def test_finance_tidak_melihat_modal_di_halaman_harga(self):
        Item.objects.filter(item_name="Palu", variant="Kecil").update(cost_amount=987654)
        self.client.force_login(self.finance)
        isi = self.client.get(reverse("master:harga")).content.decode()
        self.assertNotIn("987", isi)
        self.assertNotIn('data-kolom="cost_amount"', isi)

    def test_catatan_edit_sel_masuk_log_dan_tanda_sel(self):
        self.client.force_login(self.admin)
        info = self.info("Palu", "Kecil")
        data = self.client.post(reverse("master:barang_ubah"), json.dumps(
            {"id": info.id, "kolom": "min_price", "nilai": "9000", "catatan": "uji"}), content_type="application/json").json()
        self.assertEqual(data["tanda"], ["cost_amount", "max_price", "min_price"][1:])
        self.assertEqual(Log.objects.get(kolom="min_price").catatan, "uji")


class ModalTest(Dasar):
    def test_modal_rata_awalnya_sama_dengan_modal_terakhir(self):
        self.assertEqual(self.info("Palu", "Kecil").modal_rata, 800)
        self.assertIsNone(self.info("Kikir").modal_rata)

    def test_modal_rata_mengikuti_selama_masih_sama(self):
        info = self.info("Palu", "Kecil")
        layanan.ubah_varian(self.admin, info, "cost_amount", "900", hak=SEMUA_HAK)
        self.assertEqual(self.info("Palu", "Kecil").modal_rata, 900)
        self.assertEqual(Log.objects.filter(kolom="modal_rata", rahasia=True).count(), 1)

    def test_modal_rata_berhenti_mengikuti_setelah_dibedakan(self):
        info = self.info("Palu", "Kecil")
        layanan.ubah_varian(self.admin, info, "modal_rata", "850", hak=SEMUA_HAK)
        info.refresh_from_db()
        layanan.ubah_varian(self.admin, info, "cost_amount", "1000", hak=SEMUA_HAK)
        self.assertEqual(self.info("Palu", "Kecil").modal_rata, 850)
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").cost_amount, 1000)

    def test_finance_tidak_boleh_ubah_dan_tidak_melihat_modal_rata(self):
        with self.assertRaises(Tolak):
            layanan.ubah_varian(self.finance, self.info("Palu", "Kecil"), "modal_rata", "1", hak=PERAN["Finance"])
        VarianInfo.objects.filter(item_name="Palu", variant="Kecil").update(modal_rata=987654)
        self.client.force_login(self.finance)
        self.assertNotIn("987", self.client.get(reverse("master:harga")).content.decode())
        wb = load_workbook(BytesIO(self.client.get(reverse("master:barang") + "?export=1").content))
        self.assertNotIn("modal_rata", [c.value for c in wb["item"][1]])


class SampahTest(Dasar):
    """Keranjang sampah: buang -> hilang dari kasir; pulihkan; hapus permanen dengan password."""

    def _hak(self, user=None):
        from master.akun.akses import hak_pengguna
        return hak_pengguna(user or self.admin)

    def test_buang_hilang_dari_item_dan_bisa_dipulihkan_utuh(self):
        from master.layanan import sampah
        from master.models import Sampah
        info = self.info("Oracle", "PD 135")
        Item.objects.filter(item_name="Oracle").update(sku="OR-135", stock_alert=4)
        stok = sampah.buang(self.admin, info, hak=self._hak(), alasan="tidak dijual lagi")
        self.assertEqual(stok, 10)
        self.assertFalse(Item.objects.filter(item_name="Oracle").exists())
        self.assertTrue(VarianInfo.objects.filter(pk=info.pk).exists())
        s = Sampah.objects.get(varian=info)
        self.assertEqual((s.alasan, s.oleh), ("tidak dijual lagi", "admin"))
        layanan.pastikan_info()  # tidak membuat VarianInfo dobel / error
        self.assertEqual(VarianInfo.objects.filter(item_name="Oracle").count(), 1)
        sampah.pulihkan(self.admin, s, hak=self._hak())
        b = Item.objects.get(item_name="Oracle", variant="PD 135")
        self.assertEqual((b.sku, b.stock_alert, b.in_stock, b.min_price, b.max_price, b.cost_amount),
                         ("OR-135", 4, 10, 1500000, 135000, 700000))
        self.assertFalse(Sampah.objects.exists())
        self.assertEqual(Log.objects.filter(kolom="sampah").count(), 2)

    def test_finance_tidak_boleh(self):
        from master.layanan import sampah
        with self.assertRaises(Tolak):
            sampah.buang(self.finance, self.info("Kikir"), hak=self._hak(self.finance))

    def test_hapus_permanen_butuh_password_dan_tanpa_riwayat(self):
        from master.layanan import sampah
        from master.models import Sampah
        info = self.info("Kikir")
        sampah.buang(self.admin, info, hak=self._hak())
        s = Sampah.objects.get(varian=info)
        with self.assertRaises(Tolak):
            sampah.hapus_permanen(self.admin, s, hak=self._hak(), password="salah")
        self.assertTrue(VarianInfo.objects.filter(pk=info.pk).exists())
        sampah.hapus_permanen(self.admin, s, hak=self._hak(), password="rahasia-123")
        self.assertFalse(VarianInfo.objects.filter(pk=info.pk).exists())
        self.assertFalse(ItemInfo.objects.filter(item_name="Kikir").exists())

    def test_terjual_di_kasir_belum_bisa_dibuang(self):
        from master.layanan import sampah
        with self.assertRaises(Tolak) as e:
            sampah.buang(self.admin, self.info("Palu", "Kecil"), hak=self._hak())  # ada di nota kasir
        self.assertIn("nota kasir", str(e.exception))
        self.assertTrue(Item.objects.filter(item_name="Palu", variant="Kecil").exists())

    def test_punya_riwayat_tidak_bisa_hapus_permanen(self):
        from master.layanan import sampah
        from master.models import MutasiStok, Sampah
        info = self.info("Kikir")
        layanan.sesuaikan_stok(self.admin, info, cara="tambah", angka=2, jenis="koreksi", hak=self._hak())
        self.assertTrue(MutasiStok.objects.filter(varian=info).exists())
        sampah.buang(self.admin, info, hak=self._hak())
        with self.assertRaises(Tolak) as e:
            sampah.hapus_permanen(self.admin, Sampah.objects.get(varian=info), hak=self._hak(), password="rahasia-123")
        self.assertIn("riwayat stok", str(e.exception))
        self.assertTrue(Sampah.objects.filter(varian=info).exists())

    def test_nama_di_sampah_tidak_bisa_dipakai_ulang_atau_dobel(self):
        from master.layanan import sampah
        sampah.buang(self.admin, self.info("Kikir"), hak=self._hak())
        with self.assertRaises(Tolak):
            layanan.tambah_varian(self.admin, item_name="Kikir", variant="N/A", category="Paku",
                                  min_price=1, max_price=1, hak=self._hak())
        with self.assertRaises(Tolak):
            layanan.ganti_nama_varian(self.admin, self.info("Palu", "Besar"), "Kecil", hak=self._hak())
        sampah.buang(self.admin, self.info("Palu", "Besar"), hak=self._hak())
        with self.assertRaises(Tolak):
            layanan.ganti_nama_item(self.admin, ItemInfo.objects.get(item_name="Palu"), "Palu Baru", hak=self._hak())

    def test_pulihkan_ditolak_kalau_nama_sudah_dipakai(self):
        from master.layanan import sampah
        from master.models import Sampah
        info = self.info("Kikir")
        sampah.buang(self.admin, info, hak=self._hak())
        buat_item("Kikir")
        with self.assertRaises(Tolak):
            sampah.pulihkan(self.admin, Sampah.objects.get(varian=info), hak=self._hak())

    def test_tidak_bisa_buang_barang_di_po_terbuka(self):
        from master.layanan import pembelian as svc
        from master.layanan import sampah
        from master.models import Supplier
        sup = Supplier.objects.create(nama="PT A")
        po = svc.buat(self.admin, supplier_id=sup.id, tanggal_pesan="2026-10-01")
        svc.tambah_baris(po, varian_id=self.info("Kikir").id, jumlah=1, harga=100)
        with self.assertRaises(Tolak) as e:
            sampah.buang(self.admin, self.info("Kikir"), hak=self._hak())
        self.assertIn(po.nomor, str(e.exception))
        self.assertTrue(Item.objects.filter(item_name="Kikir").exists())

    def test_halaman(self):
        self.client.force_login(self.admin)
        info = self.info("Palu", "Besar")
        item_id = ItemInfo.objects.get(item_name="Palu").id
        r = self.client.get(reverse("master:item_panel", args=[item_id]))
        self.assertContains(r, "data-buang")
        r = self.client.post(reverse("master:item", args=[item_id]), {"aksi": "buang", "id": info.id, "alasan": "rusak"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Item.objects.filter(item_name="Palu", variant="Besar").exists())
        r = self.client.get(reverse("master:barang"))
        self.assertEqual(r.status_code, 200)
        r = self.client.get(reverse("master:sampah"))
        self.assertContains(r, "Besar")
        self.assertContains(r, "Hapus permanen")
        from master.models import Sampah
        s = Sampah.objects.get(varian=info)
        r = self.client.post(reverse("master:sampah"), {"id": s.id, "aksi": "hapus", "password": "rahasia-123"})
        self.assertFalse(VarianInfo.objects.filter(pk=info.pk).exists())
        # buang semua varian: Palu Kecil sudah terjual -> ditolak, tidak ada yang terhapus
        self.client.post(reverse("master:item", args=[item_id]), {"aksi": "buang_semua"})
        self.assertTrue(Item.objects.filter(item_name="Palu", variant="Kecil").exists())
        r = self.client.get(reverse("master:item", args=[item_id]))
        self.assertContains(r, "terjual")
        self.client.force_login(self.finance)
        self.assertEqual(self.client.get(reverse("master:sampah")).status_code, 403)


class SamakanVarianTest(Dasar):
    def test_samakan_floor_dan_jual_tanpa_modal(self):
        from master.models import ItemInfo, Log
        self.client.force_login(self.admin)
        palu = ItemInfo.objects.get(item_name="Palu")
        acuan = self.info("Palu", "Kecil")
        r = self.client.post(reverse("master:item", args=[palu.id]),
                             {"aksi": "samakan", "acuan": acuan.id, "kolom": ["min_price", "max_price"]})
        self.assertEqual(r.status_code, 302)
        besar = Item.objects.get(item_name="Palu", variant="Besar")
        self.assertEqual((besar.min_price, besar.max_price, besar.cost_amount), (1000, 1500, 2000))
        self.assertEqual(Log.objects.filter(item_name="Palu", variant="Besar", kolom__in=["min_price", "max_price"]).count(), 2)
        self.assertContains(self.client.get(reverse("master:item", args=[palu.id])), "Samakan semua varian dengan")

    def test_samakan_modal_butuh_hak_ubah_modal(self):
        from master.models import ItemInfo
        palu = ItemInfo.objects.get(item_name="Palu")
        with self.assertRaises(Tolak):
            layanan.samakan_varian(self.finance, palu, self.info("Palu", "Kecil"), ["cost_amount"],
                                   hak={"master", "ubah_harga"})
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Besar").cost_amount, 2000)

    def test_acuan_harus_varian_barang_yang_sama(self):
        from master.models import ItemInfo
        with self.assertRaises(Tolak):
            layanan.samakan_varian(self.admin, ItemInfo.objects.get(item_name="Palu"), self.info("Kikir"),
                                   ["max_price"], hak=SEMUA_HAK)


class UrutAlamiTest(Dasar):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        for v in ["12 MM", "4 MM", "8 MM", "5 MM"]:
            buat_item("Triplek", v)
        buat_item("amplas", "No 1000")
        buat_item("amplas", "No 60")
        layanan.sinkron_info()

    def test_halaman_harga_dan_stok_urut_angka_di_semua_halaman(self):
        self.client.force_login(self.admin)
        for nama in ("harga", "stok", "stok_lokasi"):
            isi = self.client.get(reverse(f"master:{nama}") + "?q=triplek").content.decode()
            posisi = [isi.index(v) for v in ["4 MM", "5 MM", "8 MM", "12 MM"]]
            self.assertEqual(posisi, sorted(posisi), nama)

    def test_export_excel_urut_alami_tanpa_peduli_huruf_besar(self):
        from master.layanan import stok_excel
        wb = load_workbook(BytesIO(excel.buat_workbook(Item.objects.all(), SEMUA_HAK)))
        baris = [(r[1].value, r[2].value) for r in wb["item"].iter_rows(min_row=2)]
        self.assertLess(baris.index(("amplas", "No 60")), baris.index(("amplas", "No 1000")))
        self.assertLess(baris.index(("amplas", "No 60")), baris.index(("Kikir", "N/A")))  # 'amplas' kecil tidak lompat ke akhir
        self.assertLess(baris.index(("Triplek", "4 MM")), baris.index(("Triplek", "12 MM")))
        wb = load_workbook(BytesIO(stok_excel.buat_workbook(Item.objects.all())))
        stok = [(r[1].value, r[2].value) for r in wb["stok"].iter_rows(min_row=2)]
        self.assertLess(stok.index(("Triplek", "5 MM")), stok.index(("Triplek", "12 MM")))


class BarangBaruPanelTest(Dasar):
    """Barang baru dari panel kanan: beberapa varian sekaligus, semua atau tidak sama sekali."""

    def _kirim(self, **data):
        import json
        dasar = {"item_name": "Taketa Martil", "category": "Alat Tukang", "brand": "Taketa", "alias": "palu, penokok",
                 "varian": [{"variant": "80 oz", "min_price": "68.000", "max_price": "70.000", "cost_amount": "60.000",
                             "in_stock": "36", "kode_rak": ""},
                            {"variant": "160 oz", "min_price": "82000", "max_price": "85000", "cost_amount": "",
                             "in_stock": "0", "kode_rak": ""}]}
        dasar.update(data)
        return self.client.post(reverse("master:barang_baru"), json.dumps(dasar), content_type="application/json")

    def test_simpan_banyak_varian_lalu_buka_di_panel(self):
        self.client.force_login(self.admin)
        r = self._kirim()
        self.assertEqual(r.status_code, 200, r.content)
        info = ItemInfo.objects.get(item_name="Taketa Martil")
        self.assertEqual(r.json()["url"], reverse("master:barang") + f"?buka={info.id}")
        v = {i.variant: (i.min_price, i.max_price, i.cost_amount, i.in_stock, i.category, i.brand)
             for i in Item.objects.filter(item_name="Taketa Martil")}
        self.assertEqual(v, {"80 oz": (68000, 70000, 60000, 36, "Alat Tukang", "Taketa"),
                             "160 oz": (82000, 85000, None, 0, "Alat Tukang", "Taketa")})
        self.assertEqual(info.alias, "palu, penokok")

    def test_satu_baris_salah_tidak_ada_yang_tersimpan(self):
        self.client.force_login(self.admin)
        r = self._kirim(varian=[{"variant": "A", "min_price": "1000", "max_price": "2000"},
                                {"variant": "B", "min_price": "3000", "max_price": "2000"}])
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["baris"], 1)
        self.assertFalse(Item.objects.filter(item_name="Taketa Martil").exists())
        r = self._kirim(varian=[{"variant": "A", "min_price": "1000", "max_price": "2000"},
                                {"variant": "a", "min_price": "1000", "max_price": "2000"}])
        self.assertEqual(r.status_code, 400)
        r = self._kirim(varian=[{"variant": "", "min_price": "1000", "max_price": "2000"},
                                {"variant": "B", "min_price": "1000", "max_price": "2000"}])
        self.assertEqual((r.status_code, r.json()["baris"]), (400, 0))
        self.assertFalse(Item.objects.filter(item_name="Taketa Martil").exists())
        # Harga boleh menyusul (27 Okt 2026): barang tersimpan tanpa harga, kasir belum bisa menjualnya.
        r = self._kirim(varian=[{"variant": "A", "min_price": "", "max_price": ""}])
        self.assertEqual(r.status_code, 200)
        self.assertTrue(Item.objects.filter(item_name="Taketa Martil", max_price__isnull=True).exists())

    def test_tanpa_varian_dan_nama_yang_sudah_ada(self):
        self.client.force_login(self.admin)
        r = self._kirim(varian=[{"variant": "", "min_price": "1000", "max_price": "2000", "in_stock": "0"}])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(list(Item.objects.filter(item_name="Taketa Martil").values_list("variant", flat=True)), ["N/A"])
        r = self._kirim(item_name="taketa martil")
        self.assertEqual(r.status_code, 400)
        self.assertIn("sudah ada", r.json()["pesan"])
        d = self.client.get(reverse("master:barang_baru") + "?cek=PALU").json()
        self.assertEqual((d["ada"], d["nama"]), (True, "Palu"))
        self.assertFalse(self.client.get(reverse("master:barang_baru") + "?cek=Obeng%20Ajaib").json()["ada"])

    def test_alamat_lama_dan_hak(self):
        self.client.force_login(self.admin)
        self.assertRedirects(self.client.get(reverse("master:barang_baru")), reverse("master:barang") + "?baru=1",
                             fetch_redirect_response=False)
        self.client.force_login(self.gudang)
        self.assertNotEqual(self._kirim().status_code, 200)
        self.assertFalse(Item.objects.filter(item_name="Taketa Martil").exists())

    def test_referensi_modal_hanya_untuk_yang_boleh(self):
        Item.objects.filter(item_name="Palu", variant="Kecil").update(cost_amount=987654)
        item_id = ItemInfo.objects.get(item_name="Palu").id
        self.client.force_login(self.admin)
        d = self.client.get(reverse("master:barang_ringkas", args=[item_id])).json()
        self.assertEqual(d["nama"], "Palu")
        self.assertIn(987654, [v.get("modal") for v in d["varian"]])
        self.client.force_login(self.finance)
        isi = self.client.get(reverse("master:barang_ringkas", args=[item_id])).content.decode()
        self.assertNotIn("987654", isi)
        self.assertNotIn("modal", isi)
