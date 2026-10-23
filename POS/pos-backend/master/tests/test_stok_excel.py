"""Uji: export dan import Excel stok (stok total dan stok per gudang)."""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from pos.models import Item

from master.akun.akses import SEMUA_HAK
from master.layanan import lokasi, stok_excel
from master.layanan.inti import Tolak
from master.models import Impor, MutasiLokasi, MutasiStok, StokLokasi
from master.tests.dasar import Dasar


class StokExcelTest(Dasar):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.toko = lokasi.lokasi_toko()
        cls.ga = lokasi.simpan_lokasi(cls.admin, hak=SEMUA_HAK, kode="ga", nama="Gudang A")
        StokLokasi.objects.create(varian=cls.info(cls, "Palu", "Kecil"), lokasi=cls.ga, jumlah=4)

    def setUp(self):
        self.client.force_login(self.admin)

    def _export(self):
        r = self.client.get(reverse("master:stok_lokasi") + "?export=1")
        self.assertEqual(r.status_code, 200)
        return load_workbook(BytesIO(r.content))

    def _cari(self, ws, nama, varian):
        for row in ws.iter_rows(min_row=2):
            if row[1].value == nama and row[2].value == varian:
                return row
        raise AssertionError("baris tidak ada")

    def _unggah(self, wb):
        keluar = BytesIO()
        wb.save(keluar)
        r = self.client.post(reverse("master:stok_impor"), {"berkas": SimpleUploadedFile(
            "stok.xlsx", keluar.getvalue())})
        return Impor.objects.order_by("-waktu").first(), r

    def test_format_export(self):
        wb = self._export()
        self.assertEqual(wb["_asal"].sheet_state, "hidden")
        kepala = [c.value for c in wb["stok"][1]]
        self.assertEqual(kepala[:6], ["id", "item_name", "variant", "category", "kode_rak", "stok_total"])
        self.assertIn("Gudang A [GA]", kepala)
        row = self._cari(wb["stok"], "Palu", "Kecil")
        i = kepala.index("Gudang A [GA]")
        self.assertEqual((row[5].value, row[i].value, row[i + 1].value), (10, 4, 6))  # total, gudang, toko

    def test_ubah_gudang_saja_berarti_pindah_dari_toko(self):
        wb = self._export()
        kepala = [c.value for c in wb["stok"][1]]
        self._cari(wb["stok"], "Palu", "Kecil")[kepala.index("Gudang A [GA]")].value = 9
        impor, r = self._unggah(wb)
        self.assertRedirects(r, reverse("master:stok_impor_pratinjau", args=[impor.id]), fetch_redirect_response=False)
        self.assertContains(self.client.get(reverse("master:stok_impor_pratinjau", args=[impor.id])), "Gudang A: 4")
        self.client.post(reverse("master:stok_impor_pratinjau", args=[impor.id]), {"aksi": "terapkan"})
        v = self.info("Palu", "Kecil")
        self.assertEqual(StokLokasi.objects.get(varian=v, lokasi=self.ga).jumlah, 9)
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Kecil").in_stock, 10)
        m = MutasiLokasi.objects.get(varian=v, sumber="impor")
        self.assertEqual((m.dari, m.ke, m.jumlah), (self.toko, self.ga, 5))

    def test_ubah_total_dan_penjualan_setelah_export_tidak_hilang(self):
        wb = self._export()
        self._cari(wb["stok"], "Palu", "Besar")[5].value = 15          # 10 -> 15 (+5)
        Item.objects.filter(item_name="Palu", variant="Besar").update(in_stock=7)  # terjual 3 setelah export
        impor, _ = self._unggah(wb)
        self.client.post(reverse("master:stok_impor_pratinjau", args=[impor.id]), {"aksi": "terapkan"})
        self.assertEqual(Item.objects.get(item_name="Palu", variant="Besar").in_stock, 12)  # 7 + 5, bukan 15
        m = MutasiStok.objects.get(varian=self.info("Palu", "Besar"), jenis="impor")
        self.assertEqual((m.jumlah, m.stok_sebelum, m.stok_sesudah), (5, 7, 12))

    def test_baris_tidak_berubah_diabaikan_dan_galat_dilaporkan(self):
        wb = self._export()
        ws = wb["stok"]
        kepala = [c.value for c in ws[1]]
        self._cari(ws, "Kikir", "N/A")[5].value = "banyak"
        self._cari(ws, "Palu", "Kecil")[kepala.index("Gudang A [GA]")].value = -1
        impor, _ = self._unggah(wb)
        self.assertEqual(impor.isi["perubahan"], [])
        self.assertEqual(len(impor.isi["galat"]), 2)

    def test_berkas_tanpa_sheet_asal_ditolak(self):
        wb = self._export()
        del wb["_asal"]
        keluar = BytesIO()
        wb.save(keluar)
        keluar.seek(0)
        with self.assertRaises(Tolak):
            stok_excel.baca_berkas(keluar, SEMUA_HAK)

    def test_akun_tanpa_hak_stok_tidak_bisa_impor(self):
        wb = self._export()
        self._cari(wb["stok"], "Palu", "Besar")[5].value = 1
        keluar = BytesIO()
        wb.save(keluar)
        keluar.seek(0)
        with self.assertRaises(Tolak):
            stok_excel.baca_berkas(keluar, {"master", "lihat_stok"})

    def test_impor_stok_tidak_muncul_di_riwayat_impor_barang(self):
        wb = self._export()
        self._cari(wb["stok"], "Palu", "Besar")[5].value = 11
        impor, _ = self._unggah(wb)
        self.assertNotContains(self.client.get(reverse("master:impor")), str(impor.id))
        self.assertRedirects(self.client.get(reverse("master:impor_pratinjau", args=[impor.id])),
                             reverse("master:stok_impor_pratinjau", args=[impor.id]), fetch_redirect_response=False)
