"""Export dan impor Excel untuk data barang.

Impor selalu dua langkah: berkas dibaca menjadi PRATINJAU (daftar perubahan),
baru diterapkan setelah disetujui. Tidak ada yang berubah saat berkas diunggah.
"""
import uuid
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from pos.models import Item

from master.layanan import inti as layanan
from master.layanan.inti import Tolak
from master.models import ItemInfo, Rak, Sampah, VarianInfo

# Urutan kolom di berkas export. `id` adalah ID tetap (VarianInfo.id).
KOLOM_EXPORT = [
    "id", "item_name", "variant", "category", "brand", "sku", "min_price", "max_price",
    "cost_amount", "modal_rata", "in_stock", "stock_alert", "kode_rak", "alias", "alias_tambahan", "aktif",
]
# Kolom yang dibaca saat impor untuk barang yang sudah ada.
KOLOM_IMPOR = [
    "category", "brand", "sku", "min_price", "max_price", "cost_amount", "modal_rata", "stock_alert",
    "kode_rak", "alias_tambahan", "aktif",
]
KOLOM_MODAL = ("cost_amount", "modal_rata")
BATAS_BARIS = 20000


def baris_export(queryset, hak):
    """Gabungkan Item dengan data master menjadi daftar dict."""
    barang = list(queryset.values(
        "item_name", "variant", "category", "brand", "sku", "min_price", "max_price",
        "cost_amount", "in_stock", "stock_alert",
    ))
    nama = {b["item_name"] for b in barang}
    info = {(v.item_name, v.variant): v for v in VarianInfo.objects.filter(item_name__in=nama)}
    alias = dict(ItemInfo.objects.filter(item_name__in=nama).values_list("item_name", "alias"))
    for b in barang:
        v = info.get((b["item_name"], b["variant"]))
        b["id"] = v.id if v else None
        b["kode_rak"] = v.rak_id if v else None
        b["alias_tambahan"] = v.alias_tambahan if v else ""
        b["aktif"] = v.aktif if v else True
        b["modal_rata"] = v.modal_rata if v else None
        b["alias"] = alias.get(b["item_name"], "")
        if "lihat_modal" not in hak:
            b.pop("cost_amount", None)
            b.pop("modal_rata", None)
    return barang


def buat_workbook(queryset, hak):
    kolom = [k for k in KOLOM_EXPORT if k not in KOLOM_MODAL or "lihat_modal" in hak]
    wb = Workbook()
    ws = wb.active
    ws.title = "item"
    ws.append(kolom)
    for b in layanan.urut_alami(baris_export(queryset, hak)):
        ws.append([("ya" if b[k] else "tidak") if k == "aktif" else b.get(k) for k in kolom])
    tebal = Font(bold=True)
    abu = PatternFill("solid", fgColor="E7E5E0")
    for sel in ws[1]:
        sel.font = tebal
        sel.fill = abu
    ws.freeze_panes = "D2"
    lebar = {"item_name": 34, "variant": 26, "category": 20, "brand": 18, "alias": 46, "alias_tambahan": 26}
    for i, k in enumerate(kolom, start=1):
        ws.column_dimensions[get_column_letter(i)].width = lebar.get(k, 13)
    ws.auto_filter.ref = ws.dimensions

    ws_rak = wb.create_sheet("rak")
    ws_rak.append(["kode_rak", "nama_rak", "lokasi", "keterangan"])
    for sel in ws_rak[1]:
        sel.font = tebal
        sel.fill = abu
    for rak in Rak.objects.all():
        ws_rak.append([rak.kode, rak.nama, rak.lokasi, rak.keterangan])
    for i, w in enumerate([12, 32, 16, 60], start=1):
        ws_rak.column_dimensions[get_column_letter(i)].width = w

    ws_info = wb.create_sheet("petunjuk")
    for baris in [
        ["Cara pakai"],
        ["cost_amount = modal terakhir. modal_rata = modal rata-rata."],
        ["1. Ubah isi sheet 'item', lalu unggah lagi lewat menu Import. Perubahan ditampilkan dulu sebelum diterapkan."],
        ["2. Kolom 'id' jangan diubah atau dihapus: itu ID tetap barang."],
        ["3. item_name dan variant tidak diubah lewat Excel. Ganti nama lewat halaman barang."],
        ["4. in_stock hanya untuk dilihat. Stok diubah lewat menu Inventory (menyusul)."],
        ["5. alias berlaku untuk semua varian satu barang; isi sama di semua barisnya. Pisahkan dengan koma."],
        ["6. Baris tanpa 'id' dan belum ada di sistem dianggap barang baru."],
    ]:
        ws_info.append(baris)
    ws_info["A1"].font = tebal
    ws_info.column_dimensions["A"].width = 110

    keluaran = BytesIO()
    wb.save(keluaran)
    return keluaran.getvalue()


# ---------------------------------------------------------------------------
# Impor: baca -> pratinjau
# ---------------------------------------------------------------------------


def _baca_sheet(ws):
    baris = ws.iter_rows(values_only=True)
    try:
        kepala = next(baris)
    except StopIteration:
        return [], []
    kepala = [layanan.rapikan_teks(k).lower() for k in kepala]
    hasil = []
    for nomor, isi in enumerate(baris, start=2):
        if isi is None or all(v is None or str(v).strip() == "" for v in isi):
            continue
        hasil.append((nomor, {k: v for k, v in zip(kepala, isi) if k}))
        if len(hasil) > BATAS_BARIS:
            raise Tolak(f"Berkas terlalu besar (lebih dari {BATAS_BARIS} baris).")
    return kepala, hasil


def _sama(lama, baru):
    return lama == baru or (lama in (None, "") and baru in (None, ""))


def baca_berkas(berkas, hak):
    """Baca berkas Excel dan kembalikan pratinjau (dict yang bisa disimpan JSON)."""
    try:
        wb = load_workbook(berkas, read_only=True, data_only=True)
    except Exception as galat:  # noqa: BLE001 - berkas dari pengguna, pesannya apa saja
        raise Tolak("Berkas tidak bisa dibuka sebagai Excel (.xlsx).") from galat

    nama_sheet = "item" if "item" in wb.sheetnames else wb.sheetnames[0]
    kepala, baris = _baca_sheet(wb[nama_sheet])
    if "item_name" not in kepala and "id" not in kepala:
        raise Tolak(f"Sheet '{nama_sheet}' tidak punya kolom 'item_name' atau 'id'.")

    # Sheet rak (opsional)
    rak_berkas = {}
    if "rak" in wb.sheetnames:
        _, baris_rak = _baca_sheet(wb["rak"])
        for _, r in baris_rak:
            kode = layanan.rapikan_teks(r.get("kode_rak") or r.get("kode"))
            if kode:
                rak_berkas[kode] = {
                    "kode": kode,
                    "nama": layanan.rapikan_teks(r.get("nama_rak") or r.get("nama")) or kode,
                    "lokasi": layanan.rapikan_teks(r.get("lokasi")),
                    "keterangan": layanan.rapikan_teks(r.get("keterangan")),
                }
    rak_ada = {r.kode: r for r in Rak.objects.all()}
    rak_perubahan = []
    for kode, r in rak_berkas.items():
        lama = rak_ada.get(kode)
        if lama is None:
            rak_perubahan.append({**r, "jenis": "baru"})
        elif (lama.nama, lama.lokasi, lama.keterangan) != (r["nama"], r["lokasi"], r["keterangan"]):
            rak_perubahan.append({**r, "jenis": "ubah"})
    rak_sah = set(rak_ada) | set(rak_berkas)

    info_per_id = {v.id: v for v in VarianInfo.objects.all()}
    info_per_kunci = {(v.item_name, v.variant): v for v in info_per_id.values()}
    barang = {
        (b["item_name"], b["variant"]): b
        for b in Item.objects.values("item_name", "variant", *layanan.KOLOM_ITEM.keys())
    }
    alias_lama = dict(ItemInfo.objects.values_list("item_name", "alias"))

    dipakai = [k for k in KOLOM_IMPOR if k in kepala]
    diabaikan = sorted(
        k for k in kepala
        if k and k not in dipakai and k not in ("id", "item_name", "variant", "alias")
    )
    dilewati_hak = []
    kolom_boleh = []
    for k in dipakai:
        if layanan.KOLOM_VARIAN[k][2] in hak:
            kolom_boleh.append(k)
        else:
            dilewati_hak.append(k)

    perubahan, baru, galat = [], [], []
    alias_berkas = {}
    sudah = set()
    for nomor, r in baris:
        info = None
        if r.get("id") not in (None, ""):
            try:
                info = info_per_id.get(int(r["id"]))
            except (TypeError, ValueError):
                info = None
            if info is None:
                galat.append({"baris": nomor, "pesan": f"id '{r.get('id')}' tidak dikenal."})
                continue
        nama = layanan.rapikan_teks(r.get("item_name"))
        varian = layanan.rapikan_varian(r.get("variant"))
        if info is None:
            if not nama:
                galat.append({"baris": nomor, "pesan": "item_name kosong."})
                continue
            info = info_per_kunci.get((nama, varian))
        elif nama and (nama, varian) != (info.item_name, info.variant) and "item_name" in kepala:
            galat.append({
                "baris": nomor,
                "pesan": f"Nama di berkas ('{nama}' / '{varian}') beda dari sistem "
                         f"('{info.item_name}' / '{info.variant}'). Nama tidak diubah; kolom lain tetap diproses.",
            })

        if info is None:
            # Barang baru
            try:
                if not {"ubah_info", "ubah_harga"} <= set(hak):
                    raise Tolak("Akun ini tidak boleh menambah barang.")
                data = {
                    "item_name": nama,
                    "variant": varian,
                    "category": layanan.rapikan_teks(r.get("category")),
                    "brand": layanan.rapikan_teks(r.get("brand")),
                    "sku": layanan.rapikan_teks(r.get("sku")) or None,
                    "min_price": layanan.baca_angka(r.get("min_price"), False),
                    "max_price": layanan.baca_angka(r.get("max_price"), False),
                    "cost_amount": layanan.baca_angka(r.get("cost_amount"), True),
                    "in_stock": layanan.baca_angka(r.get("in_stock"), True) or 0,
                    "stock_alert": layanan.baca_angka(r.get("stock_alert"), True),
                    "kode_rak": layanan.rapikan_teks(r.get("kode_rak")) or None,
                }
                if not data["category"]:
                    raise Tolak("category wajib diisi untuk barang baru.")
                if data["kode_rak"] and data["kode_rak"] not in rak_sah:
                    raise Tolak(f"Rak '{data['kode_rak']}' tidak dikenal.")
                if (nama, varian) in sudah:
                    raise Tolak("Baris kembar di berkas.")
                sudah.add((nama, varian))
                baru.append({"baris": nomor, **data})
            except Tolak as tolak:
                galat.append({"baris": nomor, "pesan": f"{nama} / {varian}: {tolak}"})
        else:
            kunci = (info.item_name, info.variant)
            if kunci in sudah:
                galat.append({"baris": nomor, "pesan": f"{kunci[0]} / {kunci[1]}: baris kembar di berkas."})
                continue
            sudah.add(kunci)
            sekarang = barang.get(kunci)
            if sekarang is None:
                ket = "ada di keranjang sampah" if Sampah.objects.filter(varian=info).exists() else "tidak ada di tabel item"
                galat.append({"baris": nomor, "pesan": f"{kunci[0]} / {kunci[1]}: {ket}."})
                continue
            sekarang = {**sekarang, "kode_rak": info.rak_id, "modal_rata": info.modal_rata,
                        "alias_tambahan": info.alias_tambahan, "aktif": info.aktif}
            for k in kolom_boleh:
                try:
                    nilai = layanan._siapkan_nilai(k, r.get(k)) if k != "kode_rak" else (
                        layanan.rapikan_teks(r.get(k)) or None
                    )
                    if k == "kode_rak" and nilai and nilai not in rak_sah:
                        raise Tolak(f"Rak '{nilai}' tidak dikenal.")
                except Tolak as tolak:
                    galat.append({"baris": nomor,
                                  "pesan": f"{kunci[0]} / {kunci[1]} - {layanan.KOLOM_VARIAN[k][0]}: {tolak}"})
                    continue
                if not _sama(sekarang.get(k), nilai):
                    perubahan.append({
                        "baris": nomor, "id": info.id, "item_name": kunci[0], "variant": kunci[1],
                        "kolom": k, "lama": sekarang.get(k), "baru": nilai,
                    })

        # Alias (tingkat item)
        if "alias" in kepala and "ubah_info" in hak:
            nama_item = info.item_name if info is not None else nama
            nilai_alias = layanan.rapikan_alias(r.get("alias"))
            if nama_item in alias_berkas and alias_berkas[nama_item]["baru"] != nilai_alias:
                if not alias_berkas[nama_item].get("bentrok"):
                    alias_berkas[nama_item]["bentrok"] = True
                    galat.append({"baris": nomor,
                                  "pesan": f"{nama_item}: alias beda antar varian di berkas. Dipakai yang pertama."})
            elif nama_item not in alias_berkas:
                alias_berkas[nama_item] = {"baru": nilai_alias}

    alias_perubahan = [
        {"item_name": n, "lama": alias_lama.get(n, ""), "baru": a["baru"]}
        for n, a in alias_berkas.items()
        if alias_lama.get(n, "") != a["baru"]
    ]

    rahasia = "lihat_modal" not in hak
    for p in perubahan:
        if p["kolom"] in KOLOM_MODAL and rahasia:
            p["lama"] = p["baru"] = None

    return {
        "sheet": nama_sheet,
        "jumlah_baris": len(baris),
        "kolom_dipakai": kolom_boleh + (["alias"] if "alias" in kepala and "ubah_info" in hak else []),
        "kolom_diabaikan": diabaikan,
        "kolom_tanpa_hak": dilewati_hak,
        "perubahan": perubahan,
        "alias": alias_perubahan,
        "baru": baru,
        "rak": rak_perubahan,
        "galat": galat,
    }


# ---------------------------------------------------------------------------
# Impor: terapkan pratinjau
# ---------------------------------------------------------------------------


def terapkan(impor, pengguna, hak, *, izinkan_baru=False):
    isi = impor.isi
    batch = uuid.uuid4()
    hasil = {"diubah": 0, "alias": 0, "baru": 0, "rak": 0, "dilewati": []}

    with transaction.atomic():
        if isi.get("rak"):
            if "kelola_rak" not in hak:
                hasil["dilewati"].append("Perubahan rak dilewati: akun ini tidak boleh mengelola rak.")
            else:
                for r in isi["rak"]:
                    _, dibuat = Rak.objects.update_or_create(
                        kode=r["kode"],
                        defaults={"nama": r["nama"], "lokasi": r["lokasi"], "keterangan": r["keterangan"]},
                    )
                    layanan.catat(pengguna, aksi="tambah" if dibuat else "ubah", sumber="impor",
                                  objek="rak", kolom="rak", baru=f"{r['kode']} - {r['nama']}", batch=batch)
                    hasil["rak"] += 1

        if isi.get("baru"):
            if not izinkan_baru:
                hasil["dilewati"].append(f"{len(isi['baru'])} barang baru dilewati (tidak dicentang).")
            else:
                for b in isi["baru"]:
                    data = {k: v for k, v in b.items() if k != "baris"}
                    try:
                        with transaction.atomic():
                            layanan.tambah_varian(pengguna, hak=hak, sumber="impor", batch=batch, **data)
                        hasil["baru"] += 1
                    except Tolak as tolak:
                        hasil["dilewati"].append(f"Baris {b['baris']}: {tolak}")

        # Perubahan diterapkan berkelompok, bukan satu per satu: impor ribuan baris
        # ke database yang jauh (mis. di cloud) jadi beberapa kueri saja.
        from master.models import Log

        info_per_id = {v.id: v for v in VarianInfo.objects.filter(
            id__in={p["id"] for p in isi.get("perubahan", [])}
        )}
        rak_sah = set(Rak.objects.values_list("kode", flat=True))
        nama = getattr(pengguna, "username", "") or "sistem"
        log_baru, info_berubah, kolom_info = [], {}, set()

        def log(item_name, variant, kolom, lama, baru):
            log_baru.append(Log(
                pengguna=pengguna, nama_pengguna=nama, aksi="ubah", sumber="impor", objek="barang",
                item_name=item_name, variant=variant, kolom=kolom, lama=layanan.teks_log(lama),
                baru=layanan.teks_log(baru), rahasia=kolom in layanan.KOLOM_RAHASIA, batch=batch,
            ))

        for p in isi.get("perubahan", []):
            info = info_per_id.get(p["id"])
            if info is None:
                hasil["dilewati"].append(f"Baris {p['baris']}: barang sudah tidak ada.")
                continue
            kolom = p["kolom"]
            try:
                if kolom in layanan.KOLOM_INFO:
                    # Kolom milik master: kumpulkan dulu, simpan sekaligus di bawah.
                    if layanan.KOLOM_VARIAN[kolom][2] not in hak:
                        raise Tolak(f"Akun ini tidak boleh mengubah {layanan.KOLOM_VARIAN[kolom][0].lower()}.")
                    if kolom == "kode_rak":
                        nilai = layanan.rapikan_teks(p["baru"]) or None
                        if nilai and nilai not in rak_sah:
                            raise Tolak(f"Rak '{nilai}' belum terdaftar.")
                        lama, atribut = info.rak_id, "rak_id"
                    else:
                        nilai = layanan._siapkan_nilai(kolom, p["baru"])
                        lama, atribut = getattr(info, kolom), kolom
                    if lama == nilai:
                        continue
                    setattr(info, atribut, nilai)
                    info_berubah[info.id] = info
                    kolom_info.add("rak" if kolom == "kode_rak" else kolom)
                    log(info.item_name, info.variant, kolom, lama, nilai)
                    hasil["diubah"] += 1
                else:
                    # Kolom di tabel item milik POS: lewat jalur biasa.
                    with transaction.atomic():
                        _, berubah = layanan.ubah_varian(
                            pengguna, info, kolom, p["baru"], hak=hak, sumber="impor", batch=batch
                        )
                    hasil["diubah"] += int(berubah)
            except Tolak as tolak:
                hasil["dilewati"].append(f"Baris {p['baris']} ({p['item_name']} / {p['variant']}): {tolak}")
        if info_berubah:
            VarianInfo.objects.bulk_update(info_berubah.values(), sorted(kolom_info), batch_size=500)

        if isi.get("alias"):
            if "ubah_info" not in hak:
                hasil["dilewati"].append("Perubahan alias dilewati: akun ini tidak boleh mengubah alias.")
            else:
                item_info = {i.item_name: i for i in ItemInfo.objects.filter(
                    item_name__in=[a["item_name"] for a in isi["alias"]]
                )}
                alias_berubah = []
                for a in isi["alias"]:
                    info = item_info.get(a["item_name"])
                    if info is None:
                        hasil["dilewati"].append(f"Alias {a['item_name']}: barang tidak ditemukan.")
                        continue
                    nilai = layanan.rapikan_alias(a["baru"])
                    if info.alias == nilai:
                        continue
                    log(info.item_name, "*", "alias", info.alias, nilai)
                    info.alias = nilai
                    alias_berubah.append(info)
                    hasil["alias"] += 1
                ItemInfo.objects.bulk_update(alias_berubah, ["alias"], batch_size=500)
        Log.objects.bulk_create(log_baru, batch_size=500)
        # Alias yang berubah lewat Excel ikut ditulis ke pencarian kasir.
        nama_alias = {a["item_name"] for a in isi.get("alias") or []}
        if "alias_tambahan" in kolom_info:
            nama_alias |= {v.item_name for v in info_berubah.values()}
        if nama_alias:
            from master.layanan import alias_kasir
            alias_kasir.tulis(nama_alias)

        impor.status = "diterapkan"
        impor.hasil = {**hasil, "batch": str(batch)}
        impor.save(update_fields=["status", "hasil"])
    return hasil
