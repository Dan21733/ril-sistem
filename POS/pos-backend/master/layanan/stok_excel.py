"""Export dan impor Excel untuk stok: stok total (kasir) dan stok per gudang.

Formatnya satu berkas yang sama untuk export dan impor, seperti di Moka: export,
ubah angkanya di Excel, lalu unggah lagi. Impor selalu lewat pratinjau.

Yang diterapkan adalah SELISIH antara angka di berkas dan angka saat export (disimpan
di sheet tersembunyi `_asal`), bukan angkanya mentah-mentah. Jadi penjualan atau barang
diambil yang terjadi di antara export dan impor tidak ikut terhapus.

- Mengubah `stok_total` = stok di kasir ikut berubah (riwayat stok, jenis "Import Excel").
- Mengubah kolom gudang dengan total tetap = barang pindah antara Toko dan gudang itu.
- Kolom Toko dihitung (total - semua gudang + dipesan), hanya untuk dilihat.
"""
import re
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from pos.models import Item

from master.layanan import inti as layanan
from master.layanan import lokasi as svc_lokasi
from master.layanan.ambil import dipesan
from master.layanan.inti import Tolak
from master.models import Lokasi, MutasiLokasi, MutasiStok, StokLokasi, VarianInfo

JENIS = "stok"
BATAS_BARIS = 20000
KODE_KOLOM = re.compile(r"\[([^\[\]]+)\]\s*$")


def _gudang():
    return list(Lokasi.objects.filter(jenis="gudang", aktif=True).order_by("urutan", "id"))


def _judul_gudang(lokasi):
    return f"{lokasi.nama} [{lokasi.kode}]"


def buat_workbook(queryset):
    """queryset: Item yang sudah disaring halaman (pencarian, kategori, rak)."""
    gudang = _gudang()
    barang = layanan.urut_alami(queryset.values("item_name", "variant", "category", "in_stock"))
    info = {(v.item_name, v.variant): v for v in VarianInfo.objects.filter(
        item_name__in={b["item_name"] for b in barang}, sampah__isnull=True).select_related("rak")}
    ids = [v.id for v in info.values()]
    per_varian = svc_lokasi.jumlah_gudang(ids)
    pesan = dipesan(ids)

    wb = Workbook()
    ws = wb.active
    ws.title = "stok"
    kolom = ["id", "item_name", "variant", "category", "kode_rak", "stok_total"] + \
        [_judul_gudang(l) for l in gudang] + ["toko (dihitung)", "dipesan (belum diambil)"]
    ws.append(kolom)
    asal = wb.create_sheet("_asal")
    asal.append(["id", "stok_total"] + [l.kode for l in gudang])
    for b in barang:
        v = info.get((b["item_name"], b["variant"]))
        if v is None:
            continue
        total = b["in_stock"] or 0
        jumlah = [per_varian.get(v.id, {}).get(l.id, 0) for l in gudang]
        tunggu = pesan.get(v.id, 0)
        ws.append([v.id, b["item_name"], b["variant"], b["category"], v.rak_id, total, *jumlah,
                   total - sum(jumlah) + tunggu, tunggu])
        asal.append([v.id, total, *jumlah])

    tebal, abu, kuning = Font(bold=True), PatternFill("solid", fgColor="E7E5E0"), PatternFill("solid", fgColor="FFF4CC")
    for sel in ws[1]:
        sel.font = tebal
        sel.fill = abu
    # Kolom yang boleh diubah ditandai kuning di judulnya.
    for i in range(6, 7 + len(gudang)):
        ws.cell(row=1, column=i).fill = kuning
    ws.freeze_panes = "D2"
    lebar = {"item_name": 34, "variant": 24, "category": 18}
    for i, k in enumerate(kolom, start=1):
        ws.column_dimensions[get_column_letter(i)].width = lebar.get(k, 16)
    ws.auto_filter.ref = ws.dimensions
    asal.sheet_state = "hidden"

    info_ws = wb.create_sheet("petunjuk")
    for baris in [
        ["Cara pakai (diekspor " + timezone.localtime().strftime("%d %b %Y %H:%M") + ")"],
        ["1. Yang boleh diubah hanya kolom berjudul kuning: stok_total dan kolom gudang."],
        ["2. stok_total = stok di kasir. Kalau diubah, stok kasir ikut berubah (tercatat di Riwayat stok)."],
        ["3. Kolom gudang = jumlah di gudang itu. Kalau diubah tanpa mengubah stok_total, barang dianggap pindah antara Toko dan gudang."],
        ["4. Kolom toko dan dipesan hanya untuk dilihat; toko dihitung dari total dikurangi semua gudang."],
        ["5. Jangan hapus kolom id dan sheet tersembunyi _asal. Yang diterapkan adalah selisih dari angka saat export,"],
        ["   jadi penjualan setelah export tidak ikut terhapus."],
        ["6. Baris yang tidak diubah diabaikan. Baris boleh dihapus kalau tidak perlu."],
        ["7. Unggah lewat tombol Import Excel di halaman Stok per gudang. Perubahan ditampilkan dulu sebelum diterapkan."],
    ]:
        info_ws.append(baris)
    info_ws["A1"].font = tebal
    info_ws.column_dimensions["A"].width = 120
    keluaran = BytesIO()
    wb.save(keluaran)
    return keluaran.getvalue()


def _angka(nilai, nomor, judul):
    if nilai is None or str(nilai).strip() == "":
        return None
    try:
        angka = float(str(nilai).strip()) if isinstance(nilai, str) else float(nilai)
    except (TypeError, ValueError):
        raise Tolak(f"Baris {nomor}, kolom {judul}: '{nilai}' bukan angka.")
    if angka != int(angka):
        raise Tolak(f"Baris {nomor}, kolom {judul}: harus bilangan bulat.")
    return int(angka)


def baca_berkas(berkas, hak):
    """Baca berkas menjadi pratinjau. Tidak ada yang berubah di sini."""
    if "ubah_stok" not in hak and "kelola_lokasi" not in hak:
        raise Tolak("Akun ini tidak boleh mengubah stok.")
    try:
        wb = load_workbook(berkas, read_only=True, data_only=True)
    except Exception:
        raise Tolak("Berkas tidak bisa dibaca. Pastikan formatnya .xlsx hasil export stok.")
    if "stok" not in wb.sheetnames or "_asal" not in wb.sheetnames:
        raise Tolak("Ini bukan berkas export stok (sheet 'stok' atau '_asal' tidak ada). Export ulang dari halaman Stok per gudang.")

    asal_rows = wb["_asal"].iter_rows(values_only=True)
    kepala_asal = [str(k or "").strip() for k in next(asal_rows, [])]
    asal = {}
    for r in asal_rows:
        if r and r[0] is not None:
            asal[int(r[0])] = dict(zip(kepala_asal, r))

    rows = wb["stok"].iter_rows(values_only=True)
    kepala = [str(k or "").strip() for k in next(rows, [])]
    if "id" not in kepala or "stok_total" not in kepala:
        raise Tolak("Kolom 'id' atau 'stok_total' tidak ditemukan.")
    gudang = {l.kode: l for l in Lokasi.objects.filter(jenis="gudang")}
    kolom_gudang = {}
    for i, judul in enumerate(kepala):
        m = KODE_KOLOM.search(judul)
        if m:
            kode = m.group(1).strip()
            lokasi = gudang.get(kode)
            if lokasi is None:
                raise Tolak(f"Gudang '{kode}' di berkas tidak ada di sistem.")
            if not lokasi.aktif:
                raise Tolak(f"Gudang {lokasi.nama} sudah tidak aktif. Export ulang.")
            kolom_gudang[i] = lokasi
    i_id, i_total = kepala.index("id"), kepala.index("stok_total")

    perubahan, galat, dibaca = [], [], 0
    info = {v.id: v for v in VarianInfo.objects.filter(pk__in=list(asal), sampah__isnull=True)}
    for nomor, isi in enumerate(rows, start=2):
        if isi is None or all(v is None or str(v).strip() == "" for v in isi):
            continue
        dibaca += 1
        if dibaca > BATAS_BARIS:
            raise Tolak(f"Berkas terlalu besar (lebih dari {BATAS_BARIS} baris).")
        try:
            vid = _angka(isi[i_id], nomor, "id")
            if vid is None or vid not in asal:
                raise Tolak(f"Baris {nomor}: id tidak dikenal. Baris baru tidak bisa ditambah lewat berkas stok.")
            v = info.get(vid)
            if v is None:
                raise Tolak(f"Baris {nomor}: barang ini sudah dihapus.")
            a = asal[vid]
            baris = {"id": vid, "nomor": nomor, "item_name": v.item_name, "variant": v.variant, "total": 0, "gudang": {}}
            total = _angka(isi[i_total], nomor, "stok_total")
            if total is not None and total != (a.get("stok_total") or 0):
                baris["total"] = total - (a.get("stok_total") or 0)
            for i, lokasi in kolom_gudang.items():
                n = _angka(isi[i] if i < len(isi) else None, nomor, lokasi.nama)
                if n is None:
                    continue
                if n < 0:
                    raise Tolak(f"Baris {nomor}, {lokasi.nama}: jumlah gudang tidak boleh minus.")
                lama = a.get(lokasi.kode) or 0
                if n != lama:
                    baris["gudang"][lokasi.kode] = n - lama
            if baris["total"] or baris["gudang"]:
                perubahan.append(baris)
        except Tolak as t:
            galat.append(str(t))
    if baris_total(perubahan) and "ubah_stok" not in hak:
        raise Tolak("Ada perubahan stok_total, tapi akun ini tidak boleh mengubah stok.")
    if any(p["gudang"] for p in perubahan) and "kelola_lokasi" not in hak:
        raise Tolak("Ada perubahan jumlah gudang, tapi akun ini tidak boleh mengelola lokasi.")
    return {"jenis": JENIS, "perubahan": perubahan, "galat": galat, "dibaca": dibaca,
            "gudang": {l.kode: l.nama for l in kolom_gudang.values()}}


def baris_total(perubahan):
    return sum(1 for p in perubahan if p["total"])


def pratinjau(isi):
    """Lengkapi pratinjau dengan angka sistem SEKARANG dan hasil akhirnya."""
    perubahan = isi.get("perubahan", [])
    ids = [p["id"] for p in perubahan]
    info = {v.id: v for v in VarianInfo.objects.filter(pk__in=ids)}
    total = {(i, v): n or 0 for i, v, n in Item.objects.filter(item_name__in={v.item_name for v in info.values()})
             .values_list("item_name", "variant", "in_stock")}
    per_varian = svc_lokasi.jumlah_gudang(ids)
    pesan = dipesan(ids)
    kode_id = dict(Lokasi.objects.filter(jenis="gudang").values_list("kode", "id"))
    nama = isi.get("gudang", {})
    hasil = []
    for p in perubahan:
        v = info.get(p["id"])
        if v is None:
            continue
        t = total.get((v.item_name, v.variant), 0)
        g = per_varian.get(v.id, {})
        semua_gudang = sum(g.values())
        toko_lama = t - semua_gudang + pesan.get(v.id, 0)
        rinci = []
        for kode, d in p["gudang"].items():
            lama = g.get(kode_id.get(kode), 0)
            rinci.append({"nama": nama.get(kode, kode), "lama": lama, "baru": lama + d, "selisih": d})
        toko_baru = toko_lama + p["total"] - sum(p["gudang"].values())
        hasil.append({**p, "total_lama": t, "total_baru": t + p["total"], "rinci": rinci,
                      "toko_lama": toko_lama, "toko_baru": toko_baru,
                      "gudang_minus": any(r["baru"] < 0 for r in rinci)})
    return hasil


def terapkan(impor, pengguna, hak):
    isi = impor.isi
    if isi.get("jenis") != JENIS:
        raise Tolak("Ini bukan impor stok.")
    if impor.status != "pratinjau":
        raise Tolak("Impor ini sudah diterapkan atau dibatalkan.")
    toko = svc_lokasi.lokasi_toko()
    gudang = {l.kode: l for l in Lokasi.objects.filter(jenis="gudang")}
    nama_pengguna = getattr(pengguna, "username", "") or "sistem"
    catatan = f"Import Excel stok ({impor.nama_berkas})"
    n_total = n_gudang = 0
    with transaction.atomic():
        for p in isi.get("perubahan", []):
            v = VarianInfo.objects.filter(pk=p["id"]).first()
            if v is None:
                continue
            baris = Item.objects.select_for_update().filter(item_name=v.item_name, variant=v.variant)
            sebelum = baris.values_list("in_stock", flat=True).first()
            if sebelum is None and not baris.exists():
                continue
            sebelum = sebelum or 0
            if p["total"]:
                if "ubah_stok" not in hak:
                    raise Tolak("Akun ini tidak boleh mengubah stok.")
                sesudah = sebelum + p["total"]
                baris.update(in_stock=sesudah)
                MutasiStok.objects.create(
                    pengguna=pengguna if getattr(pengguna, "pk", None) else None, nama_pengguna=nama_pengguna,
                    varian=v, item_name=v.item_name, variant=v.variant, jenis="impor",
                    jumlah=p["total"], stok_sebelum=sebelum, stok_sesudah=sesudah, catatan=catatan)
                n_total += 1
            for kode, d in p["gudang"].items():
                lokasi = gudang.get(kode)
                if lokasi is None or not d:
                    continue
                if "kelola_lokasi" not in hak:
                    raise Tolak("Akun ini tidak boleh mengelola lokasi.")
                s = StokLokasi.objects.select_for_update().filter(varian=v, lokasi=lokasi).first()
                if s is None:
                    s = StokLokasi.objects.create(varian=v, lokasi=lokasi, jumlah=0)
                s.jumlah += d
                s.save(update_fields=["jumlah"])
                MutasiLokasi.objects.create(
                    pengguna=pengguna if getattr(pengguna, "pk", None) else None, nama_pengguna=nama_pengguna,
                    varian=v, item_name=v.item_name, variant=v.variant,
                    dari=toko if d > 0 else lokasi, ke=lokasi if d > 0 else toko,
                    jumlah=abs(d), sumber="impor", catatan=catatan)
                n_gudang += 1
        impor.status = "diterapkan"
        impor.hasil = {"total": n_total, "gudang": n_gudang}
        impor.save(update_fields=["status", "hasil"])
        layanan.catat(pengguna, aksi="ubah", objek="stok", kolom="impor_stok",
                      baru=f"{n_total} stok total, {n_gudang} jumlah gudang", catatan=catatan)
    return impor.hasil
