"""Isi database PENGEMBANGAN dari berkas export Excel.

Membuat tabel POS (kalau belum ada) lalu mengisinya dari sheet berkas export.
Perintah ini menolak jalan kalau tabel `item` sudah berisi, supaya tidak pernah
menimpa database yang sedang dipakai. Jangan diarahkan ke database produksi.
"""
import uuid
from datetime import datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from master.layanan import inti as layanan
from master.layanan.legacy import buat_tabel_pos, tabel_ada
from master.models import ItemInfo, Rak, VarianInfo
from pos.models import (Device, DeviceTransactionSequence, Item, PaymentProof,
                        TransactionDetail, Transactions, Worker)


def _baris(wb, nama):
    if nama not in wb.sheetnames:
        return []
    data = wb[nama].iter_rows(values_only=True)
    kepala = [str(k).strip() if k is not None else "" for k in next(data)]
    return [dict(zip(kepala, isi)) for isi in data if isi and any(v is not None for v in isi)]


def _bulat(nilai):
    return None if nilai is None or nilai == "" else int(nilai)


def _waktu(nilai):
    if nilai is None:
        return None
    if not isinstance(nilai, datetime):
        nilai = datetime.fromisoformat(str(nilai))
    return timezone.make_aware(nilai, timezone.get_default_timezone()) if timezone.is_naive(nilai) else nilai


class Command(BaseCommand):
    help = "Buat tabel POS dan isi dari export Excel (hanya untuk database pengembangan)."

    def add_arguments(self, parser):
        parser.add_argument("berkas", help="Berkas export .xlsx (sheet: item, worker, transactions, ...)")

    def handle(self, *args, **options):
        buat_tabel_pos()
        if not tabel_ada("master_varian_info"):
            raise CommandError("Tabel master belum ada. Jalankan dulu: python manage.py migrate --settings=master.settings")
        if Item.objects.exists():
            raise CommandError("Tabel item sudah berisi. Perintah ini hanya untuk database kosong.")

        wb = load_workbook(options["berkas"], read_only=True, data_only=True)
        with transaction.atomic():
            barang, alias, rak_varian = [], {}, {}
            for r in _baris(wb, "item"):
                nama = layanan.rapikan_teks(r.get("item_name"))
                varian = layanan.rapikan_varian(r.get("variant"))
                barang.append(Item(
                    item_name=nama, variant=varian,
                    moka_id=None if r.get("moka_id") is None else str(r["moka_id"]),
                    category=r.get("category"), sku=r.get("sku"), brand=r.get("brand"),
                    ecommerce_item=bool(r.get("ecommerce_item")),
                    min_price=_bulat(r.get("min_price")), max_price=_bulat(r.get("max_price")),
                    in_stock=_bulat(r.get("in_stock")), track_stock=bool(r.get("track_stock")),
                    track_alert=bool(r.get("track_alert")), stock_alert=_bulat(r.get("stock_alert")),
                    cost_amount=_bulat(r.get("cost_amount")),
                ))
                if r.get("alias") and nama not in alias:
                    alias[nama] = layanan.rapikan_alias(r["alias"])
                if r.get("kode_rak"):
                    rak_varian[(nama, varian)] = str(r["kode_rak"]).strip()
            Item.objects.bulk_create(barang, batch_size=500)

            for r in _baris(wb, "rak"):
                Rak.objects.update_or_create(kode=str(r["kode_rak"]).strip(), defaults={
                    "nama": r.get("nama_rak") or r["kode_rak"], "lokasi": r.get("lokasi") or "",
                    "keterangan": r.get("keterangan") or "",
                })
            for kode in set(rak_varian.values()) - set(Rak.objects.values_list("kode", flat=True)):
                Rak.objects.create(kode=kode, nama=kode)

            layanan.sinkron_info()
            for info in ItemInfo.objects.filter(item_name__in=alias):
                info.alias = alias[info.item_name]
                info.save(update_fields=["alias"])
            daftar = list(VarianInfo.objects.all())
            for info in daftar:
                info.rak_id = rak_varian.get((info.item_name, info.variant))
            VarianInfo.objects.bulk_update(daftar, ["rak"], batch_size=500)

            Worker.objects.bulk_create([Worker(id=r["id"], name=r["name"], pin=r["pin"]) for r in _baris(wb, "worker")])
            Device.objects.bulk_create([Device(id=uuid.UUID(r["id"]), device_code=r["device_code"]) for r in _baris(wb, "device")])
            DeviceTransactionSequence.objects.bulk_create([
                DeviceTransactionSequence(device_id=uuid.UUID(r["device_id"]), next_sequence=_bulat(r["next_sequence"]),
                                          updated_at=_waktu(r["updated_at"]))
                for r in _baris(wb, "device_transaction_sequence")
            ])
            Transactions.objects.bulk_create([
                Transactions(
                    sync_id=uuid.UUID(r["sync_id"]), transaction_id=r.get("transaction_id"), worker_id=r["worker_id"],
                    transaction_time=_waktu(r["transaction_time"]), item_amount=_bulat(r["item_amount"]),
                    shipping_amount=_bulat(r.get("shipping_amount")), total_amount=_bulat(r["total_amount"]),
                    payment_method=r["payment_method"], payment_amount=_bulat(r["payment_amount"]),
                    change_amount=_bulat(r["change_amount"]),
                ) for r in _baris(wb, "transactions")
            ])
            TransactionDetail.objects.bulk_create([
                TransactionDetail(
                    sync_id=uuid.UUID(r["sync_id"]), item_name=layanan.rapikan_teks(r["item_name"]),
                    variant=layanan.rapikan_varian(r.get("variant")), quantity=_bulat(r["quantity"]),
                    unit_price=_bulat(r["unit_price"]), discount=_bulat(r.get("discount")) or 0,
                    subtotal=_bulat(r["subtotal"]),
                ) for r in _baris(wb, "transaction_detail")
            ])
            PaymentProof.objects.bulk_create([
                PaymentProof(
                    proof_id=uuid.UUID(r["proof_id"]), sync_id=uuid.UUID(r["sync_id"]), object_key=r["object_key"],
                    content_type=r["content_type"], byte_size=_bulat(r["byte_size"]), sha256=r["sha256"],
                    captured_at=_waktu(r.get("captured_at")), uploaded_at=_waktu(r["uploaded_at"]),
                    device_code=r.get("device_code"),
                ) for r in _baris(wb, "payment_proof")
            ])

        self.stdout.write(self.style.SUCCESS(
            f"Terisi: {Item.objects.count()} varian, {ItemInfo.objects.count()} barang, "
            f"{Rak.objects.count()} rak, {Transactions.objects.count()} nota."
        ))
