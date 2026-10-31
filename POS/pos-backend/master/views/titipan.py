"""Setoran barang titipan: penjualan barang milik toko titipan, dikurangi potongan toko, disetor ke pemiliknya."""
from datetime import timedelta
from io import BytesIO
from urllib.parse import quote

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from master.akun.akses import perlu
from master.layanan import titipan as svc
from master.layanan.inti import Tolak
from master.layanan.pembelian import baca_tanggal
from master.models import SetoranTitipan, Toko


def _rp(n):
    return "Rp" + f"{int(n or 0):,}".replace(",", ".")


def _periode(request, toko):
    """Tanggal dari/sampai di alamat; default sejak setoran terakhir (atau awal bulan) sampai hari ini."""
    hari_ini = timezone.localdate()
    awal = svc.awal_periode(toko) or hari_ini.replace(day=1)
    try:
        dari = baca_tanggal(request.GET.get("dari")) or awal
        sampai = baca_tanggal(request.GET.get("sampai")) or hari_ini
    except Tolak:
        dari, sampai = awal, hari_ini
    return dari, min(sampai, hari_ini)


def _teks_wa(h):
    t = h["total"]
    baris = [f"*Laporan penjualan {h['toko'].nama}*", f"{h['dari']:%d-%m-%Y} s.d. {h['sampai']:%d-%m-%Y}", ""]
    for r in h["rincian"][:40]:
        nama = r["item_name"] + ("" if r["variant"] in ("", "N/A") else f" / {r['variant']}")
        qty = r["qty"] - r["qty_retur"]
        baris.append(f"- {nama}: {qty} = {_rp(r['bersih'])}")
    if len(h["rincian"]) > 40:
        baris.append(f"... dan {len(h['rincian']) - 40} barang lain (lihat Excel)")
    baris += ["", f"Penjualan: {_rp(t['jual'])}"]
    if t["retur"]:
        baris.append(f"Retur: -{_rp(t['retur'])}")
    baris += [f"Potongan toko: -{_rp(t['potongan'])}", f"*Disetor: {_rp(t['disetor'])}*"]
    return "\n".join(baris)


def _excel(toko, dari, sampai, rincian, total, judul="Setoran"):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "setoran"
    ws.append([f"{judul} {toko.nama}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{dari:%d-%m-%Y} s.d. {sampai:%d-%m-%Y}"])
    ws.append([])
    kepala = ["Barang", "Varian", "Terjual", "Penjualan", "Retur (qty)", "Retur (Rp)", "Bersih", "Potongan %",
              "Potongan toko", "Disetor"]
    ws.append(kepala)
    for sel in ws[4]:
        sel.font = Font(bold=True)
    for r in rincian:
        ws.append([r["item_name"], "" if r["variant"] == "N/A" else r["variant"], r["qty"], r["jual"], r["qty_retur"],
                   r["retur"], r["jual"] - r["retur"], r.get("persen", ""), r["potongan"], r["disetor"]])
    ws.append([])
    ws.append(["Total", "", total.get("qty", ""), total["jual"], total.get("qty_retur", ""), total["retur"],
               total["jual"] - total["retur"], "", total["potongan"], total["disetor"]])
    for sel in ws[ws.max_row]:
        sel.font = Font(bold=True)
    for kol, lebar in zip("ABCDEFGHIJ", (34, 22, 9, 14, 11, 12, 14, 12, 14, 14)):
        ws.column_dimensions[kol].width = lebar
    isi = BytesIO()
    wb.save(isi)
    respons = HttpResponse(isi.getvalue(),
                           content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respons["Content-Disposition"] = f'attachment; filename="setoran_{toko.kode}_{dari:%Y%m%d}_{sampai:%Y%m%d}.xlsx"'
    return respons


@perlu("lihat_modal")
def daftar(request):
    """Semua toko titipan: berapa yang belum disetor sejak setoran terakhir."""
    hari_ini = timezone.localdate()
    baris = []
    for t in svc.daftar_toko():
        dari = svc.awal_periode(t) or hari_ini.replace(day=1)
        h = svc.hitung(t, dari, hari_ini) if dari <= hari_ini else None
        baris.append({"toko": t, "dari": dari, "h": h, "terakhir": svc.setoran_terakhir(t)})
    if len(baris) == 1:
        return redirect("master:titipan_toko", toko_id=baris[0]["toko"].pk)
    return render(request, "master/titipan_daftar.html", {"menu": "titipan", "baris": baris})


@perlu("lihat_modal")
def toko(request, toko_id):
    t = get_object_or_404(Toko, pk=toko_id, titip_di__isnull=False)
    if request.method == "POST":
        try:
            dari = baca_tanggal(request.POST.get("dari"), wajib=True, label="Tanggal awal")
            sampai = baca_tanggal(request.POST.get("sampai"), wajib=True, label="Tanggal akhir")
            s = svc.tandai_disetor(request.user, t, dari, sampai, request.POST.get("catatan", ""))
            messages.success(request, f"Setoran {_rp(s.disetor)} tercatat. Periode berikutnya mulai {s.sampai + timedelta(days=1):%d-%m-%Y}.")
        except Tolak as tolak:
            messages.error(request, str(tolak))
            return redirect(f"{reverse('master:titipan_toko', args=[t.pk])}?dari={request.POST.get('dari', '')}&sampai={request.POST.get('sampai', '')}")
        return redirect("master:titipan_toko", toko_id=t.pk)

    dari, sampai = _periode(request, t)
    h = svc.hitung(t, dari, sampai)
    if request.GET.get("excel") == "1":
        return _excel(t, dari, sampai, h["rincian"], h["total"], judul="Penjualan barang titipan")
    terakhir = svc.setoran_terakhir(t)
    telepon = "".join(c for c in (t.telepon or "") if c.isdigit())
    if telepon.startswith("0"):
        telepon = "62" + telepon[1:]
    return render(request, "master/titipan_toko.html", {
        "menu": "titipan", "t": t, "h": h, "dari": dari, "sampai": sampai, "terakhir": terakhir,
        "boleh_setor": not terakhir or dari > terakhir.sampai,
        "wa": f"https://wa.me/{telepon}?text={quote(_teks_wa(h))}" if h["rincian"] else "",
        "riwayat": SetoranTitipan.objects.filter(toko=t)[:24],
    })


@perlu("lihat_modal")
def setoran(request, setoran_id):
    s = get_object_or_404(SetoranTitipan.objects.select_related("toko"), pk=setoran_id)
    total = {"qty": sum(r["qty"] for r in s.rincian), "qty_retur": sum(r["qty_retur"] for r in s.rincian),
             "jual": s.total_jual, "retur": s.total_retur, "potongan": s.potongan, "disetor": s.disetor}
    if request.GET.get("excel") == "1":
        return _excel(s.toko, s.dari, s.sampai, s.rincian, total)
    for r in s.rincian:
        r["bersih"] = r["jual"] - r["retur"]
    return render(request, "master/titipan_setoran.html", {"menu": "titipan", "s": s, "total": total})
